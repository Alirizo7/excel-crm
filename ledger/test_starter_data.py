import hashlib
from io import BytesIO
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .models import (
    Delivery,
    ImportBatch,
    Operation,
    PartnerBalance,
    SourceSheet,
)


class StarterWorkspaceTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls.media = TemporaryDirectory()
        cls.media_override = override_settings(MEDIA_ROOT=cls.media.name)
        cls.media_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls.media_override.disable()
        cls.media.cleanup()

    def test_bootstrap_installs_exact_snapshot_once_and_keeps_existing_data(self):
        output = StringIO()
        call_command('bootstrap_workspace', stdout=output)
        self.assertIn('Демо-данные установлены', output.getvalue())
        self.assertEqual(ImportBatch.objects.count(), 1)
        self.assertEqual(SourceSheet.objects.count(), 13)
        self.assertEqual(Operation.objects.count(), 1765)
        self.assertEqual(Delivery.objects.count(), 634)
        self.assertEqual(PartnerBalance.objects.count(), 46)
        batch = ImportBatch.objects.get()
        self.assertEqual(
            hashlib.sha256(Path(batch.file.path).read_bytes()).hexdigest(),
            batch.sha256,
        )
        user = get_user_model().objects.get(username='admin')
        self.assertTrue(user.check_password('admin123'))
        self.assertTrue(self.client.login(username='admin', password='admin123'))
        for kind in ('operations', 'deliveries', 'partners', 'report'):
            response = self.client.get(reverse('export_download', args=[kind]))
            self.assertEqual(response.status_code, 200)
            exported = load_workbook(BytesIO(response.content), read_only=True)
            self.assertTrue(exported.sheetnames)
            exported.close()
        user.first_name = 'Не перезаписывать'
        user.save(update_fields=['first_name'])
        output = StringIO()
        call_command('bootstrap_workspace', stdout=output)
        self.assertIn('существующая база сохранена', output.getvalue())
        user.refresh_from_db()
        self.assertEqual(user.first_name, 'Не перезаписывать')
        self.assertEqual(Operation.objects.count(), 1765)
        self.assertEqual(ImportBatch.objects.count(), 1)

    def test_starter_assets_are_present_and_match_the_import(self):
        seed = Path(settings.BASE_DIR) / 'starter_data'
        workbook = seed / '01.09.2026.xlsx'
        fixture = seed / 'workspace.json.gz'
        self.assertTrue(fixture.is_file())
        self.assertEqual(
            hashlib.sha256(workbook.read_bytes()).hexdigest(),
            'fd3183a8287e8bcd6e698cffd40f689d949d1725a3fcdf9977f5068154897df5',
        )
