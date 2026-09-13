import copy
import hashlib
import json
import math
import os
import uuid
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import skipUnless
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .models import WorkingWorkbook, Operation, Delivery, PartnerBalance, DebtPayment, RecordChange
from .tests import legacy_bytes
from .services.changes import snapshot, changed, log_change
from .services.debts import record_payment
from .services.importing import stage_import, commit_import
from .services.workbook_files import read_xlsx, validate_snapshot, calculate, write_xlsx, formula_errors, cells
from .services.workbooks import create_book, save_book, restore_book, preview_book, apply_book, WorkbookConflict


class WorkbookTests(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.get(username='admin'))

    @classmethod
    def setUpClass(cls):
        cls.media = TemporaryDirectory()
        cls.media_override = override_settings(MEDIA_ROOT=cls.media.name)
        cls.media_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls.media_override.disable(); cls.media.cleanup()

    @classmethod
    def setUpTestData(cls):
        cls.batch, _ = stage_import(legacy_bytes(), '01.09.2026.xlsx')
        commit_import(cls.batch.pk)
        cls.book = create_book(cls.batch.pk, 'tester')

    def edited(self, sheet, r, c, value, formula=False, allowed=True):
        self.book.refresh_from_db()
        data = copy.deepcopy(self.book.data)
        ws = next(s for s in data['sheets'].values() if s['name'] == sheet)
        cell = ws['cellData'].setdefault(str(r-1), {}).setdefault(str(c-1), {})
        cell.pop('f', None); cell.pop('v', None)
        if value is not None: cell['f' if formula else 'v'] = value
        self.book = save_book(self.book.pk, data, self.book.revision, 'tester', allowed)
        return self.book

    def apply(self):
        self.book.refresh_from_db()
        plans, blocked, token, *_ = preview_book(self.book)
        self.assertEqual(blocked, [])
        self.assertFalse([p['errors'] for p in plans if p['errors']])
        return apply_book(self.book.pk, self.book.revision, token, 'tester')

    def test_create_is_idempotent_and_preserves_source_and_records(self):
        initial = Path(self.batch.file.path).read_bytes()
        self.assertEqual(create_book(self.batch.pk,'other').pk, self.book.pk)
        self.assertEqual(WorkingWorkbook.objects.count(),1)
        self.assertEqual(Operation.objects.count(),4)
        self.assertEqual(formula_errors(self.book.data),[])
        self.assertEqual(self.apply(),0)
        self.assertEqual(Path(self.batch.file.path).read_bytes(),initial)

    def test_live_formula_chain_save_export_restore_and_stale_write(self):
        initial = self.book.revision
        self.edited('тура ака',4,6,30000,allowed=False)
        ws=next(s for s in self.book.data['sheets'].values() if s['name']=='тура ака')
        self.assertEqual(ws['cellData']['3']['7']['v'],16110)
        self.assertEqual(ws['cellData']['3']['11']['v'],59010930)
        self.assertEqual(Delivery.objects.get().gross,26610)
        payload=write_xlsx(self.book.data)
        cached=load_workbook(BytesIO(payload),data_only=True)
        formulas=load_workbook(BytesIO(payload),data_only=False)
        self.assertEqual(cached['тура ака']['L4'].value,59010930)
        self.assertEqual(formulas['тура ака']['L4'].value,'=J4*K4')
        with self.assertRaises(WorkbookConflict):save_book(self.book.pk,self.book.data,initial,'tester',True)
        self.assertEqual(self.apply(),1)
        self.assertEqual(Delivery.objects.get().amount,Decimal('59010930'))
        self.book=restore_book(self.book.pk,initial,self.book.revision,'tester',True)
        self.assertEqual(self.apply(),1)
        self.assertEqual(Delivery.objects.get().gross,26610)
        self.assertEqual(self.book.versions.count(),3)

    def test_formula_permissions_and_literal_text(self):
        with self.assertRaises(ValidationError):self.edited('тура ака',4,8,'=1+1',True,False)
        self.edited('Приход-Расход',4,3,'=Это текст',False,False)
        ws=next(s for s in self.book.data['sheets'].values() if s['name']=='Приход-Расход')
        self.assertNotIn('f',ws['cellData']['3']['2'])
        out=load_workbook(BytesIO(write_xlsx(self.book.data)),data_only=False)
        self.assertEqual(out['Приход-Расход']['C4'].data_type,'s')

    def test_apply_preserves_payments_and_rejects_changed_preview(self):
        debt=PartnerBalance.objects.get(name='Партнёр А')
        record_payment(debt.pk,amount=Decimal('100'),date=date(2026,9,1),note='',request_key=uuid.uuid4(),revision=debt.revision,actor='tester')
        self.edited('Кунлик',5,4,1500)
        plans,blocked,token,*_=preview_book(self.book)
        self.assertEqual(blocked,[])
        debt.refresh_from_db()
        record_payment(debt.pk,amount=Decimal('50'),date=date(2026,9,1),note='',request_key=uuid.uuid4(),revision=debt.revision,actor='tester')
        with self.assertRaises(WorkbookConflict):apply_book(self.book.pk,self.book.revision,token,'tester')
        self.assertEqual(self.apply(),1)
        debt.refresh_from_db();self.assertEqual(debt.balance,1250);self.assertEqual(debt.paid_amount,150)
        self.assertEqual(debt.outstanding,1100);self.assertEqual(DebtPayment.objects.count(),2)
        self.assertEqual(self.apply(),0)
        self.edited('Кунлик',5,4,300)
        plans,*_=preview_book(self.book);self.assertTrue(any(p['errors'] for p in plans))

    def test_unrelated_crm_edit_preserved_and_same_field_conflict_blocked(self):
        record=Operation.objects.get(kind='income',source_row=4)
        record.description='Правка CRM';record.save()
        self.edited('Приход-Расход',4,5,900)
        self.assertEqual(self.apply(),1)
        record.refresh_from_db();self.assertEqual(record.description,'Правка CRM');self.assertEqual(record.amount,900)
        self.edited('Приход-Расход',4,5,1000)
        record.amount=950;record.save()
        plans,*_=preview_book(self.book);self.assertTrue(plans[0]['errors'])

    def test_invalid_number_does_not_silently_delete_operation(self):
        self.edited('Приход-Расход',4,5,'не число')
        plans,blocked,*_=preview_book(self.book)
        self.assertTrue(blocked)
        self.assertEqual(Operation.objects.filter(active=True).count(),4)

    def test_soft_delete_and_restore_keep_uid_and_history(self):
        record=Operation.objects.get(kind='income',source_row=4);uid=record.uid
        self.edited('Приход-Расход',4,5,None)
        self.assertEqual(self.apply(),1)
        record.refresh_from_db();self.assertFalse(record.active)
        self.book=restore_book(self.book.pk,1,self.book.revision,'tester',True)
        self.assertEqual(self.apply(),1)
        record.refresh_from_db();self.assertTrue(record.active);self.assertEqual(record.uid,uid)
        self.assertEqual(Operation.objects.count(),4)
        self.assertTrue(RecordChange.objects.filter(record_id=record.pk,action='restore').exists())

    def test_new_formula_errors_and_inconsistent_delivery_totals_block_apply(self):
        self.edited('Приход-Расход',4,5,'=1/0',True)
        plans,blocked,*_=preview_book(self.book);self.assertTrue(blocked)
        self.book=restore_book(self.book.pk,1,self.book.revision,'tester',True)
        self.edited('тура ака',4,12,'=J4*K4+100',True)
        self.assertContains(self.client.get(reverse('workbook_history',args=[self.book.pk])),'=J4*K4+100')
        plans,blocked,*_=preview_book(self.book)
        self.assertTrue(any(p['errors'] for p in plans))

    def test_invalid_structures_external_formulas_and_resources(self):
        for change in [lambda d: d['sheets']['sheet-0'].update(columnCount=100000),
                       lambda d: d['sheets']['sheet-0']['rowData']['3'].update(h='bad'),
                       lambda d: d['sheets']['sheet-0'].update(mergeData=[{'startRow':-1}]),
                       lambda d: d['styles'].update(evil={'fs':float('inf')}),
                       lambda d: d['sheets']['sheet-0']['cellData']['3']['1'].update(f='=WEBSERVICE("https://example.com")'),
                       lambda d: d['sheets']['sheet-0']['cellData']['3']['1'].update(f="='[other.xlsx]Sheet1'!A1")]:
            with self.subTest(change=change):
                data=copy.deepcopy(self.book.data);change(data)
                with self.assertRaises(ValidationError):validate_snapshot(data,self.book.data)
        data=copy.deepcopy(self.book.data);data['resources']=[{'name':'external','data':'x'}]
        self.assertNotIn('resources',validate_snapshot(data,self.book.data))

    def test_pages_post_validation_and_archive(self):
        for url in ['workbooks','workbook_detail','workbook_history','workbook_preview','workbook_data','workbook_download']:
            response=self.client.get(reverse(url,args=[] if url=='workbooks' else [self.book.pk]),follow=True)
            self.assertEqual(response.status_code,200,url)
        response=self.client.post(reverse('workbook_save',args=[self.book.pk]),data='{}',content_type='application/json')
        self.assertEqual(response.status_code,400)
        new,_=stage_import(legacy_bytes(income=999),'02.09.2026.xlsx');commit_import(new.pk)
        with self.assertRaises(ValidationError):self.edited('Приход-Расход',4,5,777)
        self.assertContains(self.client.get(reverse('workbook_detail',args=[self.book.pk])),'data-readonly="true"')

    def test_anonymous_cannot_read_financial_workbook(self):
        self.client.logout()
        response=self.client.get(reverse('workbook_data',args=[self.book.pk]))
        self.assertEqual(response.status_code,302)
        user=get_user_model().objects.create_user('operator',password='test-only-password')
        self.client.force_login(user)
        self.assertContains(self.client.get(reverse('workbook_detail',args=[self.book.pk])),'data-formulas="false"')

    def test_default_admin_can_edit_formulas(self):
        self.assertContains(self.client.get(reverse('workbook_detail', args=[self.book.pk])), 'data-formulas="true"')

    def finish(self, **extra):
        return self.client.post(reverse('workbook_finish', args=[self.book.pk]),
            data=json.dumps({'revision': self.book.revision, **extra}), content_type='application/json')

    def test_home_opens_current_table_and_import_opens_next_table(self):
        self.assertRedirects(self.client.get(reverse('home')), reverse('workbook_detail', args=[self.book.pk]))
        new, _ = stage_import(legacy_bytes(income=999), '02.09.2026.xlsx')
        response = self.client.post(reverse('import_confirm', args=[new.pk]), {'debt_policy': 'keep'})
        new_book = WorkingWorkbook.objects.get(batch=new)
        self.assertRedirects(response, reverse('workbook_detail', args=[new_book.pk]))
        self.assertRedirects(self.client.get(reverse('home')), reverse('workbook_detail', args=[new_book.pk]))
        self.book.batch.refresh_from_db()
        self.assertEqual(self.book.batch.status, 'superseded')

    def test_finish_updates_formula_results_and_retry_does_not_duplicate(self):
        self.edited('тура ака', 4, 6, 30000, allowed=False)
        response = self.finish()
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['count'], 1)
        delivery = Delivery.objects.get()
        self.assertEqual(delivery.gross, 30000)
        self.assertEqual(delivery.amount, Decimal('59010930'))
        self.book.refresh_from_db()
        self.assertEqual(self.book.applied_revision, self.book.revision)
        self.assertEqual(self.finish().json()['count'], 0)
        self.assertEqual(Delivery.objects.count(), 1)

    def test_finish_projects_only_once(self):
        from .services.workbooks import project
        self.edited('Приход-Расход', 4, 5, 900)
        with patch('ledger.services.workbooks.project', wraps=project) as projection:
            response = self.finish()
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(projection.call_count, 1)
        self.assertEqual(Operation.objects.get(kind='income', source_row=4).amount, 900)

    def test_prepared_projection_rejects_a_new_workbook_revision(self):
        self.edited('Приход-Расход', 4, 5, 900)
        _, _, token, projected, issues = preview_book(self.book)
        revision = self.book.revision
        self.edited('Приход-Расход', 4, 5, 1000)
        with self.assertRaises(WorkbookConflict):
            apply_book(self.book.pk, revision, token, 'tester', prepared=(revision, projected, issues))
        with self.assertRaises(WorkbookConflict):
            apply_book(self.book.pk, self.book.revision, token, 'tester', prepared=(revision, projected, issues))
        self.assertNotEqual(Operation.objects.get(kind='income', source_row=4).amount, 900)

    def test_prepared_projection_rechecks_crm_payment_under_lock(self):
        self.edited('Кунлик', 5, 4, 1500)
        _, _, token, projected, issues = preview_book(self.book)
        debt = PartnerBalance.objects.get(name='Партнёр А')
        record_payment(debt.pk, amount=Decimal('100'), date=date(2026,9,1), note='',
            request_key=uuid.uuid4(), revision=debt.revision, actor='tester')
        with self.assertRaises(WorkbookConflict):
            apply_book(self.book.pk, self.book.revision, token, 'tester',
                prepared=(self.book.revision, projected, issues))
        debt.refresh_from_db()
        self.assertEqual(debt.paid_amount, 100)
        self.assertNotEqual(debt.balance, 1250)

    def test_finish_invalid_values_block_all_changes_and_keep_copy(self):
        self.edited('тура ака', 4, 6, 30000)
        self.edited('Приход-Расход', 4, 5, 'не число')
        response = self.finish()
        self.assertEqual(response.status_code, 422)
        self.assertTrue(response.json()['issues'])
        self.assertEqual(Delivery.objects.get().gross, 26610)
        self.assertEqual(Operation.objects.filter(active=True).count(), 4)
        self.book.refresh_from_db()
        self.assertGreater(self.book.revision, self.book.applied_revision)

    def test_finish_requires_explicit_confirmation_for_removals(self):
        record = Operation.objects.get(kind='income', source_row=4)
        self.edited('Приход-Расход', 4, 5, None)
        response = self.finish()
        self.assertTrue(response.json()['confirmation_required'])
        self.assertTrue(response.json()['removals'])
        record.refresh_from_db()
        self.assertTrue(record.active)
        response = self.finish(confirmation_token=response.json()['token'])
        self.assertEqual(response.status_code, 200, response.content)
        record.refresh_from_db()
        self.assertFalse(record.active)
        self.assertTrue(RecordChange.objects.filter(record_id=record.pk, action='delete').exists())

    def test_finish_rejects_changed_record_after_removal_confirmation(self):
        record = Operation.objects.get(kind='income', source_row=4)
        self.edited('Приход-Расход', 4, 5, None)
        confirmation = self.finish().json()['token']
        record.amount += 10
        record.save()
        response = self.finish(confirmation_token=confirmation)
        self.assertIn(response.status_code, (409, 422))
        record.refresh_from_db()
        self.assertTrue(record.active)

    def test_finish_preserves_payments_and_blocks_debt_below_paid_amount(self):
        debt = PartnerBalance.objects.get(name='Партнёр А')
        record_payment(debt.pk, amount=Decimal('100'), date=date(2026, 9, 1), note='',
            request_key=uuid.uuid4(), revision=debt.revision, actor='tester')
        self.edited('Кунлик', 5, 4, 1500)
        self.assertEqual(self.finish().status_code, 200)
        debt.refresh_from_db()
        self.assertEqual(debt.paid_amount, 100)
        self.assertEqual(debt.outstanding, 1150)
        self.edited('Кунлик', 5, 4, 300)
        self.assertEqual(self.finish().status_code, 422)
        debt.refresh_from_db()
        self.assertEqual(debt.outstanding, 1150)

    def test_finish_checks_revision_method_csrf_and_archived_books(self):
        url = reverse('workbook_finish', args=[self.book.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        secure = Client(enforce_csrf_checks=True)
        secure.force_login(get_user_model().objects.get(username='admin'))
        self.assertEqual(secure.post(url, data='{}', content_type='application/json').status_code, 403)
        self.assertEqual(self.finish(revision=-1).status_code, 409)
        self.assertEqual(self.client.post(url, data='{}', content_type='application/json').status_code, 400)
        self.edited('Приход-Расход', 4, 5, 777)
        new, _ = stage_import(legacy_bytes(income=999), '02.09.2026.xlsx')
        commit_import(new.pk)
        self.assertEqual(self.finish().status_code, 400)

    def test_removed_sections_redirect_to_relevant_pages(self):
        self.assertRedirects(self.client.get(reverse('exports') + '?start=2026-09-01'),
            reverse('dashboard') + '?start=2026-09-01#reports')
        self.assertRedirects(self.client.get(reverse('sheets')), reverse('imports') + '#file-history')
        self.assertRedirects(self.client.get(reverse('help')), reverse('workbooks'), fetch_redirect_response=False)


@skipUnless(os.getenv('WORKBOOK_PATH'), 'Set WORKBOOK_PATH to check the original customer workbook')
class OriginalWorkbookParityTests(TestCase):
    def test_all_formulas_preserved_except_explicit_legacy_range_error(self):
        original=Path(os.environ['WORKBOOK_PATH']).read_bytes();checksum=hashlib.sha256(original).hexdigest()
        data=calculate(validate_snapshot(read_xlsx(original,'original.xlsx')))
        cached=load_workbook(BytesIO(original),data_only=True)
        count=0
        for sid,r,c,cell in cells(data):
            if not cell.get('f'):continue
            count+=1;name=data['sheets'][sid]['name'];expected=cached[name].cell(r+1,c+1).value
            if name=='тура ака' and (r,c)==(877,0):self.assertEqual(cell['v'],'#N/A');continue
            if isinstance(expected, datetime):
                from openpyxl.utils.datetime import to_excel
                expected=to_excel(expected)
            if isinstance(expected,(int,float)):
                self.assertTrue(math.isclose(cell.get('v',0),expected,rel_tol=1e-12,abs_tol=1e-5),(name,r+1,c+1,expected,cell))
            else:self.assertEqual(cell.get('v'),expected,(name,r+1,c+1))
        self.assertEqual(count,13104)
        self.assertEqual(len(formula_errors(data)),2)
        # Verify the exported formulas and literal values, not only cached results.
        exported = load_workbook(BytesIO(write_xlsx(data)), data_only=False)
        source = load_workbook(BytesIO(original), data_only=False)
        self.assertEqual(exported.sheetnames, source.sheetnames)
        for source_sheet in source:
            for row in source_sheet:
                for cell in row:
                    if cell.value is None: continue
                    actual = exported[source_sheet.title][cell.coordinate]
                    expected = '=' + cell.value[2:] if cell.data_type == 'f' and cell.value.startswith('=+') else cell.value
                    self.assertEqual(actual.value, expected, (source_sheet.title, cell.coordinate))
                    self.assertEqual(actual.number_format, cell.number_format, (source_sheet.title, cell.coordinate))
        self.assertEqual(hashlib.sha256(Path(os.environ['WORKBOOK_PATH']).read_bytes()).hexdigest(),checksum)
