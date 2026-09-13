import copy
from datetime import datetime
from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment

from .services.workbook_files import read_xlsx, validate_snapshot, calculate, write_xlsx


class FormulaEngineTests(SimpleTestCase):
    def test_supported_functions_cross_sheet_errors_and_recalculation(self):
        workbook = Workbook()
        ws = workbook.active; ws.title = 'Ввод'
        for row, value in enumerate([10, 0, -2.5, 'текст', None, True], 1):
            ws.cell(row, 1, value)
        formulas = {
            'A1': ('=SUM(Ввод!A1:A6)', 7.5),
            'A2': ('=MIN(Ввод!A1:A6)', -2.5),
            'A3': ('=MAX(Ввод!A1:A6)', 10),
            'A4': ('=AVERAGE(Ввод!A1:A6)', 2.5),
            'A5': ('=COUNT(Ввод!A1:A6)', 3),
            'A6': ('=COUNTA(Ввод!A1:A6)', 5),
            'A7': ('=IF(Ввод!A1>0,"Да","Нет")', 'Да'),
            'A8': ('=IFERROR(1/0,9)', 9),
            'A9': ('=ROUND(2.675,2)', 2.68),
            'A10': ('=ROUNDUP(-2.34,1)', -2.4),
            'A11': ('=ROUNDDOWN(-2.34,1)', -2.3),
            'A12': ('=ABS(Ввод!A3)', 2.5),
            'A13': ('=NA()', '#N/A'),
            'A14': ('=1/0', '#DIV/0!'),
            'A15': ('=Ввод!A4*2', '#VALUE!'),
            'A16': ('=Ввод!A5+1', 1),
            'A17': ('=2^3+4*5-6/2', 25),
            'A18': ('=SUM(A1,A3)*50%', 8.75),
            'A19': ('=Ввод!$A$1+Ввод!A$3', 7.5),
        }
        result = workbook.create_sheet('Итоги')
        for address, (formula, _) in formulas.items(): result[address] = formula
        output = BytesIO(); workbook.save(output)
        data = calculate(validate_snapshot(read_xlsx(output.getvalue(), 'functions.xlsx')))
        exported = load_workbook(BytesIO(write_xlsx(data)), data_only=True)['Итоги']
        for address, (_, expected) in formulas.items():
            with self.subTest(address=address): self.assertEqual(exported[address].value, expected)
        data['sheets']['sheet-0']['cellData']['0']['0']['v'] = 3.5
        data = calculate(validate_snapshot(data))
        exported = load_workbook(BytesIO(write_xlsx(data)), data_only=True)['Итоги']
        self.assertEqual(exported['A1'].value, 1)
        self.assertEqual(exported['A18'].value, 2.25)
        self.assertEqual(exported['A19'].value, 1)

    def test_export_cached_styles_do_not_share_mutable_merge_borders(self):
        workbook = Workbook(); ws = workbook.active
        for address in ['A1', 'C1']:
            ws[address] = 12.5
            ws[address].font = Font(name='Arial', size=12, bold=True, color='123456')
            ws[address].fill = PatternFill('solid', fgColor='ABCDEF')
            ws[address].alignment = Alignment(horizontal='right')
            ws[address].number_format = '#,##0.00'
        ws['B1'] = 0
        ws['B1'].border = Border(right=Side(style='thin', color='FF0000'))
        ws['A2'] = datetime(2026,9,13)
        other = workbook.create_sheet('Другой'); other['A1'] = 12.5
        other['A1']._style = copy.copy(ws['A1']._style)
        output = BytesIO(); workbook.save(output)
        data = read_xlsx(output.getvalue(), 'styles.xlsx')
        data['sheets']['sheet-0']['mergeData'] = [dict(startRow=0,endRow=0,startColumn=0,endColumn=1)]
        exported = load_workbook(BytesIO(write_xlsx(data)), data_only=False)
        self.assertEqual(exported.worksheets[0]['A1'].border.right.style, 'thin')
        for cell in [exported.worksheets[0]['C1'], exported['Другой']['A1']]:
            self.assertFalse(cell.border.right and cell.border.right.style)
            self.assertEqual(cell.font.color.rgb[-6:], '123456')
            self.assertEqual(cell.fill.fgColor.rgb[-6:], 'ABCDEF')
            self.assertEqual(cell.number_format, '#,##0.00')
            self.assertTrue(cell.font.bold)
        self.assertEqual(exported.worksheets[0]['A2'].value, datetime(2026,9,13))
