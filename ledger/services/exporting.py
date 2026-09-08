"""Native Python Excel exports used by the deployed Django application."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from django.utils.translation import override
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from .importing import OP_HEADERS, DEL_HEADERS


def table(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B2'
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='142D2A')
        cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for values in rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # User-supplied strings must remain text, including leading =, + and @.
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=11, color='213732')
            cell.alignment = Alignment(vertical='center')
            if sheet.max_row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F1F7F4')
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = 'dd.mm.yyyy'
            elif isinstance(cell.value, (Decimal, float, int)):
                cell.number_format = '#,##0.00'
        sheet.row_dimensions[sheet.max_row].height = 23
    for i, header in enumerate(headers, 1):
        width = 22
        if any(x in header for x in ('Назначение', 'Примечание', 'Контрагент', 'Сообщение')):
            width = 48
        if header == 'ID':
            width = 39
        if header == 'Дата':
            width = 15
        sheet.column_dimensions[get_column_letter(i)].width = max(width, min(len(header)+3, 34))
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


@override('ru')
def make_export(kind, operations=(), deliveries=(), partners=(), issues=(), template=False):
    # Exchange sheet names, columns and choice values form one stable contract.
    # Changing the interface language must not break export/reimport.
    workbook = Workbook()
    workbook.remove(workbook.active)
    if kind in ('operations', 'report'):
        table(workbook, 'Операции', OP_HEADERS, [] if template else
              [[str(o.uid), o.date, o.get_kind_display(), o.amount, o.description, o.category, o.partner] for o in operations])
    if kind in ('deliveries', 'report'):
        headers = DEL_HEADERS + ([] if template else ['Нетто кг', 'Чистый вес кг', 'Стоимость UZS'])
        table(workbook, 'Поставки', headers, [] if template else
              [[str(d.uid), d.date, d.get_direction_display(), d.partner, d.vehicle, d.gross, d.tare,
                d.discount, d.price, d.notes, d.net, d.clean_weight, d.amount] for d in deliveries])
    if kind in ('partners', 'report'):
        table(workbook, 'Взаиморасчёты', ['Контрагент', 'Дебет исходный UZS', 'Кредит исходный UZS', 'Сальдо UZS', 'Примечание'],
              [[p.name, p.receivable_column, p.payable_column, p.balance, p.note] for p in partners])
    if kind == 'report':
        table(workbook, 'Проверка данных', ['Уровень', 'Лист', 'Ячейка', 'Сообщение'],
              [[i['level'], i['sheet'], i['cell'], i['message']] for i in issues])
    if template:
        sheet = workbook.active
        values = 'Приход,Расход' if kind == 'operations' else 'Отгрузка,Приёмка'
        validation = DataValidation(type='list', formula1=f'"{values}"', allow_blank=False)
        validation.error = 'Выберите значение из списка.'
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add('C2:C12000')
        sheet.column_dimensions['B'].number_format = 'dd.mm.yyyy'
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
