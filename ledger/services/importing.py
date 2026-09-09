from django.utils.translation import gettext_noop
"""Adapters for the supplied workbook and the application's exchange templates.

Workbook cells are data, never instructions. Formula evaluation is deliberately
restricted; cached Excel results retain provenance and unsupported cells are reported.
"""
import ast
import hashlib
import json
import operator
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import openpyxl
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl.utils import get_column_letter
from ledger.models import Activity, Delivery, ImportBatch, Operation, PartnerBalance, SourceSheet
from ledger.services.changes import workspace_transaction, partner_for

MAX_BYTES = 10 * 1024 * 1024
ROLES = {'тонн': gettext_noop('Объёмы и остатки'), 'Хитой': gettext_noop('Расчёты и распределение прибыли'),
         'тура ака': gettext_noop('Отгрузки металла и взаиморасчёты'), 'дафтарча': gettext_noop('Приёмка по контрагентам'),
         'Кунлик': gettext_noop('Сводка взаиморасчётов'), 'Приход-Расход': gettext_noop('Основной денежный журнал'),
         'Чермт Бонус': gettext_noop('Бонусы по чермету'), 'Рахим': gettext_noop('Расчёты с Рахимом'), 'Каракул': gettext_noop('Объёмы по объектам'),
         'тогом': gettext_noop('Оплаты и стоимость товара'), 'Гостиница': gettext_noop('Расходы по гостинице'),
         'Бонусча': gettext_noop('Дополнительные бонусы'), 'Чугун': gettext_noop('Приёмка чугуна')}
OP_HEADERS = ['ID', gettext_noop('Дата'), gettext_noop('Тип'), 'Сумма UZS', gettext_noop('Назначение'), gettext_noop('Категория'), gettext_noop('Контрагент')]
DEL_HEADERS = ['ID', gettext_noop('Дата'), gettext_noop('Направление'), gettext_noop('Контрагент'), 'Машина', 'Брутто кг', 'Тара кг', 'Скидка %', 'Цена UZS/кг', gettext_noop('Примечание')]


def serial(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, Decimal):
        return str(value)
    return value


def num(value):
    if value is None or isinstance(value, bool) or (isinstance(value, str) and not value.strip()):
        return None
    try:
        result = Decimal(str(value).replace('\xa0', '').replace(' ', '').replace(',', '.'))
        return result if result.is_finite() and abs(result) < Decimal('1e17') else None
    except InvalidOperation:
        return None


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        parts = re.split(r'[,./\s-]+', value.strip())
        if len(parts) == 3 and all(p.isdigit() for p in parts):
            a, b, c = map(int, parts)
            try:
                return date(a, b, c) if len(parts[0]) == 4 else date(c + 2000 if c < 100 else c, b, a)
            except ValueError:
                pass
    return None


class CellReader:
    def __init__(self, formulas, cached, issues):
        self.formulas, self.cached, self.issues = formulas, cached, issues
        self.memo, self.visiting = {}, set()

    def get(self, sheet, row, col):
        key = (sheet, row, col)
        if key in self.memo:
            return self.memo[key]
        cell = self.formulas[sheet].cell(row, col)
        value = self.cached[sheet].cell(row, col).value
        if cell.data_type == 'f' and value is None:
            if key in self.visiting or len(self.visiting) > 80:
                raise ValueError('Циклическая формула')
            self.visiting.add(key)
            try:
                value = self.formula(sheet, cell.value)
            except (ValueError, TypeError, ArithmeticError, SyntaxError, KeyError, RecursionError):
                value = None
                self.issues.append({'level': 'warning', 'sheet': sheet, 'cell': cell.coordinate,
                                    'message': gettext_noop('Нет сохранённого результата формулы. Пересчитайте файл в Excel; ячейка недоступна.')})
            finally:
                self.visiting.discard(key)
        if isinstance(value, str) and value.startswith(('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!')):
            value = None
        self.memo[key] = value
        return value

    def formula(self, sheet, expression):
        expression = expression.lstrip('=+')
        if len(expression) > 2000 or '[' in expression:
            raise ValueError('Unsupported formula')
        ref = r"(?:(?:'([^']+)'|([\w ]+))!)?(\$?[A-Z]{1,3}\$?\d+)"

        def cell_value(match):
            target = match.group(1) or match.group(2) or sheet
            cell = self.formulas[target][match.group(3).replace('$', '')]
            if cell.row > 12000 or cell.column > 64:
                raise ValueError('Range too large')
            return self.get(target, cell.row, cell.column)

        single = re.fullmatch(ref, expression)
        if single:
            return cell_value(single)

        def sum_range(match):
            first, last = match.group(1).replace('$', ''), match.group(2).replace('$', '')
            a, b = self.formulas[sheet][first], self.formulas[sheet][last]
            if b.row > 12000 or b.column > 64 or (b.row-a.row+1)*(b.column-a.column+1)>12000:
                raise ValueError('Range too large')
            total = Decimal(0)
            for row in range(a.row, b.row + 1):
                for col in range(a.column, b.column + 1):
                    value = self.get(sheet, row, col)
                    source = self.formulas[sheet].cell(row, col)
                    if value is None and source.data_type == 'f':
                        raise ValueError('Unavailable dependency')
                    total += num(value) or 0
            return str(total)

        expression = re.sub(r'SUM\((\$?[A-Z]+\$?\d+):(\$?[A-Z]+\$?\d+)\)', sum_range, expression, flags=re.I)

        def replace_ref(match):
            value = cell_value(match)
            if value is None:
                target = match.group(1) or match.group(2) or sheet
                if self.formulas[target][match.group(3).replace('$', '')].value is not None:
                    raise ValueError('Unavailable dependency')
            return str(num(value) or 0)

        expression = re.sub(ref, replace_ref, expression)
        tree = ast.parse(expression, mode='eval')
        allowed = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv}

        def visit(node):
            if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                return Decimal(str(node.value))
            if isinstance(node, ast.BinOp) and type(node.op) in allowed:
                return allowed[type(node.op)](visit(node.left), visit(node.right))
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                return visit(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
            raise ValueError('Unsupported formula')
        return visit(tree.body)


def category(description):
    t = str(description).lower()
    for name, words in [(gettext_noop('Металл'), ('чермет', 'металл', 'чугун', 'релс', 'мяхки', 'новои', 'турон')),
                        (gettext_noop('Транспорт'), ('мошин', 'машин', 'йул', 'хура', 'бензин')),
                        (gettext_noop('Расчёты с партнёрами'), ('рахим', 'тура', 'ойбек', 'хитой', 'нурбек', 'фарход')),
                        (gettext_noop('Хозяйственные расходы'), ('бозор', 'ужин', 'уйга', 'рас', 'усто'))]:
        if any(word in t for word in words):
            return name
    return gettext_noop('Прочее')


def inspect_workbook(data, filename, report_date=None):
    if len(data) > MAX_BYTES or not filename.lower().endswith('.xlsx'):
        raise ValidationError(gettext_noop('Выберите файл .xlsx размером не больше 10 МБ.'))
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 60 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise ValidationError(gettext_noop('Файл слишком велик после распаковки.'))
            if any('vbaProject' in i.filename for i in archive.infolist()):
                raise ValidationError(gettext_noop('Файлы с макросами не поддерживаются.'))
        formulas = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(BytesIO(data), data_only=True, keep_links=False)
    except (BadZipFile, OSError, KeyError, ValueError, SyntaxError) as exc:
        raise ValidationError(gettext_noop('Не удалось прочитать Excel. Проверьте, что файл не повреждён и не защищён паролем.')) from exc
    if (len(formulas.sheetnames) > 30 or any(s.max_row > 12000 or s.max_column > 64 for s in formulas)
            or sum(s.max_row*s.max_column for s in formulas) > 750000):
        raise ValidationError(gettext_noop('Поддерживаются книги до 30 листов, 12 000 строк и 64 колонок на лист.'))
    issues, sheets = [], []
    reader = CellReader(formulas, cached, issues)
    operations, deliveries, balances = [], [], []
    report_date = report_date or parse_date(Path(filename).stem)
    for position, sheet in enumerate(formulas):
        rows, formula_count = [], 0
        for cells in sheet:
            populated = []
            useful = False
            for cell in cells:
                if cell.value is None:
                    populated.append({'value': None})
                    continue
                raw = cell.value
                value = cached[sheet.title][cell.coordinate].value
                if cell.data_type == 'f':
                    formula_count += 1
                if (cell.data_type != 'f' and str(raw).strip()) or (cell.data_type == 'f' and value not in (None, 0, '')):
                    useful = True
                if isinstance(value, str) and value.startswith(('#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!')):
                    issues.append({'level': 'warning', 'sheet': sheet.title, 'cell': cell.coordinate,
                                   'message': f'Ошибка Excel {value}. Исходная формула сохранена; значение исключено из расчётов.'})
                populated.append({'value': serial(value), 'formula': raw if cell.data_type == 'f' else '',
                                  'original': serial(raw) if cell.data_type != 'f' else None})
            if useful:
                rows.append({'number': cells[0].row, 'cells': populated})
        sheets.append({'name': sheet.title, 'position': position, 'rows': rows, 'columns': sheet.max_column,
                       'formula_count': formula_count, 'role': ROLES.get(sheet.title, gettext_noop('Исходные данные'))})

    if {'Приход-Расход', 'тура ака', 'Кунлик'}.issubset(formulas.sheetnames):
        fmt = 'legacy'
        checks = [('Приход-Расход', 'A1', gettext_noop('Дата')), ('Приход-Расход', 'B1', 'Расход сумма'),
                  ('Приход-Расход', 'D1', gettext_noop('Дата')), ('Приход-Расход', 'E1', 'Приход сумма'),
                  ('тура ака', 'F1', gettext_noop('Брутто')), ('тура ака', 'G1', gettext_noop('Тара')), ('тура ака', 'I1', '% скидка'),
                  ('тура ака', 'K1', gettext_noop('Цена')), ('Кунлик', 'D1', 'А ман карз'), ('Кунлик', 'E1', 'Ман Карз')]
        if any(str(formulas[s][c].value).strip() != expected for s, c, expected in checks):
            raise ValidationError(gettext_noop('Названия листов знакомы, но колонки отличаются от образца. Импорт остановлен, чтобы не перепутать суммы.'))
        for amount_col, date_col, desc_col, kind in [(2, 1, 3, 'expense'), (5, 4, 6, 'income')]:
            current_date = None
            for row in range(4, formulas['Приход-Расход'].max_row + 1):
                raw_date = reader.get('Приход-Расход', row, date_col)
                if raw_date is not None and str(raw_date).strip():
                    current_date = parse_date(raw_date)
                value = num(reader.get('Приход-Расход', row, amount_col))
                if value is None and formulas['Приход-Расход'].cell(row, amount_col).data_type == 'f':
                    issues.append({'level': 'error', 'sheet': 'Приход-Расход', 'cell': f'{get_column_letter(amount_col)}{row}',
                                   'message': gettext_noop('Не удалось прочитать сумму операции. Пересчитайте и исправьте формулу перед импортом.')})
                if value is None or value == 0:
                    continue
                description = str(reader.get('Приход-Расход', row, desc_col) or 'Без назначения').strip()
                operations.append({'date': serial(current_date), 'kind': kind,
                                   'amount': str(value.quantize(Decimal('.01'))),
                                   'description': description[:500],
                                   'category': category(description), 'partner': '', 'source_sheet': 'Приход-Расход', 'source_row': row})
        for name, mapping, direction in [('тура ака', (1, 5, 6, 7, 9, 11, 4), 'out'),
                                         ('дафтарча', (2, 5, 6, 7, 9, 10, 3), 'in'),
                                         ('Чугун', (3, 7, 8, 9, 11, 13, 4), 'in')]:
            if name not in formulas.sheetnames:
                continue
            current_date, partner = None, ('Тура ака' if name == 'тура ака' else name)
            dc, vc, gc, tc, sc, pc, nc = mapping
            for row in range(1, formulas[name].max_row + 1):
                if name == 'дафтарча' and str(reader.get(name, row, 2)).strip() == gettext_noop('Дата'):
                    partner, current_date = str(reader.get(name, row, 11) or 'Без названия'), None
                raw_date = reader.get(name, row, dc)
                if raw_date is not None and str(raw_date).strip():
                    current_date = parse_date(raw_date)
                gross, tare = num(reader.get(name, row, gc)), num(reader.get(name, row, tc))
                if row < 4 or (gross is None and tare is None):
                    continue
                # Summary cells and repeated section totals are not deliveries.
                if formulas[name].cell(row, gc).data_type == 'f' and 'SUM' in str(formulas[name].cell(row, gc).value).upper():
                    continue
                discount, price = num(reader.get(name, row, sc)), num(reader.get(name, row, pc))
                if gross is None or tare is None or gross < tare or tare < 0 or gross >= Decimal('1e11') or price is None or not 0 <= price < Decimal('1e12') or not 0 <= (discount or 0) <= 100:
                    issues.append({'level': 'warning', 'sheet': name, 'cell': f'{get_column_letter(gc)}{row}',
                                   'message': gettext_noop('Поставка пропущена: проверьте брутто, тару, цену и скидку.')})
                    continue
                if gross == 0 and tare == 0:
                    continue
                net, clean, amount = Delivery.calculate(gross, tare, discount or 0, price)
                if amount >= Decimal('1e18'):
                    issues.append({'level': 'error', 'sheet': name, 'cell': f'{get_column_letter(pc)}{row}',
                                   'message': gettext_noop('Стоимость поставки превышает допустимый размер числа.')})
                    continue
                deliveries.append({'date': serial(current_date), 'direction': direction, 'partner': partner[:150],
                                   'vehicle': str(reader.get(name, row, vc) or 'Без номера')[:150],
                                   'gross': str(gross), 'tare': str(tare), 'discount': str(discount or 0), 'price': str(price),
                                   'net': str(net), 'clean_weight': str(clean), 'amount': str(amount),
                                   'notes': str(reader.get(name, row, nc) or '')[:500], 'source_sheet': name, 'source_row': row})
        for row in range(5, formulas['Кунлик'].max_row + 1):
            name = reader.get('Кунлик', row, 2)
            debit, credit = num(reader.get('Кунлик', row, 4)), num(reader.get('Кунлик', row, 5))
            if name is None or (not debit and not credit):
                continue
            debit, credit = debit or Decimal(0), credit or Decimal(0)
            balances.append({'name': str(name).strip()[:250], 'receivable_column': str(debit.quantize(Decimal('.01'))),
                             'payable_column': str(credit.quantize(Decimal('.01'))), 'balance': str((debit-credit).quantize(Decimal('.01'))),
                             'note': str(reader.get('Кунлик', row, 3) or '')[:500], 'source_sheet': 'Кунлик', 'source_row': row})
        if 'Гостиница' in formulas.sheetnames:
            issues.append({'level': 'info', 'sheet': 'Гостиница', 'cell': 'A1:C23',
                           'message': gettext_noop('Смешанные суммы UZS/USD и текстовые заметки сохранены в исходном листе, без пересчёта валют.')})
        total_formula = str(formulas['тура ака']['L3'].value or '')
        adjustment = re.search(r'\)\s*\+\s*(\d+(?:\.\d+)?)\s*$', total_formula)
        if adjustment:
            issues.append({'level': 'info', 'sheet': 'тура ака', 'cell': 'L3',
                           'message': f'В исходный итог включена корректировка {Decimal(adjustment.group(1)):,.2f} UZS. Она не является поставкой и не добавляется к стоимости отгрузок.'})
        opening_label = str(reader.get('Приход-Расход', 4, 3) or '')
        if 'карз' in opening_label.lower() or 'долг' in opening_label.lower():
            issues.append({'level': 'info', 'sheet': 'Приход-Расход', 'cell': 'B4',
                           'message': gettext_noop('Начальный долг из строки 4 сохранён как расход по исходному журналу. Сальдо журнала не означает прибыль.')})
    elif len(formulas.sheetnames) == 1 and formulas.active.title in (gettext_noop('Операции'), gettext_noop('Поставки')):
        fmt = 'operations' if formulas.active.title == gettext_noop('Операции') else 'deliveries'
        sheet = formulas.active
        headers = OP_HEADERS if fmt == 'operations' else DEL_HEADERS
        if [sheet.cell(1, i+1).value for i in range(len(headers))] != headers:
            raise ValidationError(gettext_noop('Колонки не совпадают с шаблоном. Скачайте шаблон на странице импорта.'))
        seen_ids = set()
        for row in range(2, sheet.max_row+1):
            values = [reader.get(sheet.title, row, i+1) for i in range(len(headers))]
            if all(v is None for v in values):
                continue
            try:
                uid = str(uuid.UUID(str(values[0]))) if values[0] else str(uuid.uuid5(uuid.NAMESPACE_URL, json.dumps([fmt, row, values[1:]], default=serial, ensure_ascii=False)))
                if uid in seen_ids:
                    raise ValueError(gettext_noop('ID повторяется внутри файла'))
                seen_ids.add(uid)
                dt = parse_date(values[1])
                if values[1] is not None and not dt:
                    raise ValueError(gettext_noop('некорректная дата'))
                if fmt == 'operations':
                    kind = {gettext_noop('Приход'): 'income', gettext_noop('Расход'): 'expense'}.get(values[2])
                    amount = num(values[3])
                    if kind is None or amount is None or amount == 0 or not values[4]:
                        raise ValueError(gettext_noop('нужны тип, ненулевая сумма и назначение'))
                    record = {'uid': uid, 'date': serial(dt), 'kind': kind, 'amount': str(amount), 'description': str(values[4]),
                              'category': str(values[5] or gettext_noop('Прочее')), 'partner': str(values[6] or '')}
                    model = Operation
                else:
                    direction = {gettext_noop('Отгрузка'): 'out', gettext_noop('Приёмка'): 'in'}.get(values[2])
                    gross, tare, discount, price = [num(v) for v in values[5:9]]
                    if direction is None or not values[3] or not values[4] or any(v is None for v in (gross, tare, discount, price)):
                        raise ValueError(gettext_noop('заполните направление, контрагента, машину и числовые поля'))
                    if not (gross >= tare >= 0 and 0 <= discount <= 100 and price >= 0):
                        raise ValueError(gettext_noop('недопустимые веса, цена или скидка'))
                    net, clean, amount = Delivery.calculate(gross, tare, discount, price)
                    record = {'uid': uid, 'date': serial(dt), 'direction': direction, 'partner': str(values[3]), 'vehicle': str(values[4]),
                              'gross': str(gross), 'tare': str(tare), 'discount': str(discount), 'price': str(price),
                              'notes': str(values[9] or ''), 'net': str(net), 'clean_weight': str(clean), 'amount': str(amount)}
                    model = Delivery
                record.update(source_sheet=sheet.title, source_row=row)
                obj = model(**record)
                obj.full_clean(exclude=['uid', 'batch', 'source_row'], validate_unique=False, validate_constraints=False)
                if model.objects.filter(uid=uid).exists():
                    issues.append({'level': 'info', 'sheet': sheet.title, 'cell': f'A{row}',
                                   'message': gettext_noop('ID уже существует. Строка пропущена; существующая запись не изменяется.')})
                else:
                    (operations if fmt == 'operations' else deliveries).append(record)
            except (ValueError, InvalidOperation, ValidationError) as exc:
                issues.append({'level': 'error', 'sheet': sheet.title, 'cell': f'A{row}', 'message': f'Строка {row}: {exc}'[:450]})
    else:
        raise ValidationError(gettext_noop('Формат не распознан. Загрузите полную книгу с листами «Приход-Расход», «тура ака», «Кунлик» либо наш шаблон «Операции» / «Поставки».'))
    undated = sum(not r['date'] for r in operations)
    if undated:
        issues.append({'level': 'warning', 'sheet': 'Приход-Расход' if fmt == 'legacy' else gettext_noop('Операции'), 'cell': '',
                       'message': f'{undated} денежных операций без даты. Они входят в общий итог, но не в отчёты за период.'})
    summary = {'operations': len(operations), 'deliveries': len(deliveries), 'partners': len(balances),
               'sheets': len(sheets), 'undated': undated, 'errors': sum(i['level']=='error' for i in issues),
               'warnings': sum(i['level']=='warning' for i in issues),
               'income': str(sum(Decimal(x['amount']) for x in operations if x['kind']=='income')),
               'expense': str(sum(Decimal(x['amount']) for x in operations if x['kind']=='expense'))}
    return {'format': fmt, 'report_date': serial(report_date), 'operations': operations, 'deliveries': deliveries,
            'balances': balances, 'sheets': sheets, 'issues': issues, 'summary': summary}


def stage_import(data, filename, report_date=None):
    digest = hashlib.sha256(data).hexdigest()
    existing = ImportBatch.objects.filter(sha256=digest).first()
    if existing:
        return existing, False
    result = inspect_workbook(data, filename, report_date)
    with transaction.atomic():
        batch = ImportBatch(filename=Path(filename).name[:255], sha256=digest, format=result['format'],
                            report_date=result['report_date'], summary=result['summary'], issues=result['issues'],
                            payload={k: result[k] for k in ('operations', 'deliveries', 'balances')})
        batch.file.save(f'{uuid.uuid4().hex}.xlsx', ContentFile(data), save=False)
        batch.save()
        SourceSheet.objects.bulk_create([SourceSheet(batch=batch, **sheet) for sheet in result['sheets']])
    return batch, True


def import_impact():
    edited = sum(model.objects.filter(batch__format='legacy', batch__status='imported', modified_at__isnull=False).count()
                 for model in (Operation, Delivery))
    balances = PartnerBalance.objects.filter(batch__format='legacy', archived_at=None)
    return {'record_edits': edited, 'debt_edits': balances.filter(modified_at__isnull=False).count(),
            'has_debts': balances.exists()}


@workspace_transaction()
def commit_import(batch_id, *, debt_policy=None, replace_local_changes=False):
    batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
    if batch.status != 'preview':
        raise ValidationError(gettext_noop('Этот файл уже обработан. Повторный импорт не выполнен.'))
    if batch.summary.get('errors'):
        raise ValidationError(gettext_noop('Исправьте ошибки в файле и загрузите его заново. Частичный импорт отключён.'))
    if batch.format == 'legacy':
        impact = import_impact()
        if debt_policy not in (None, 'keep', 'replace'):
            raise ValidationError(gettext_noop('Выберите способ обновления долгов.'))
        if impact['debt_edits'] and debt_policy is None:
            raise ValidationError(gettext_noop('Есть изменения долгов или погашения. Выберите, сохранить текущие долги или принять остатки из новой книги.'))
        debt_policy = debt_policy or 'replace'
        if (impact['record_edits'] or (impact['debt_edits'] and debt_policy == 'replace')) and not replace_local_changes:
            raise ValidationError(gettext_noop('Подтвердите замену изменений импортированных записей новым снимком.'))
        previous = ImportBatch.objects.filter(format='legacy', status='imported')
        for model in (Operation, Delivery):
            model.objects.filter(batch__in=previous).update(active=False, archived_at=timezone.now())
        if debt_policy == 'replace':
            PartnerBalance.objects.filter(batch__format='legacy', archived_at=None).update(active=False, archived_at=timezone.now())
        previous.update(status='superseded')
    for name, model in [('operations', Operation), ('deliveries', Delivery), ('balances', PartnerBalance)]:
        if name == 'balances' and debt_policy == 'keep':
            continue
        records = batch.payload.get(name, [])
        if name == 'balances':
            records = [dict(record, counterparty=partner_for(record['name'])) for record in records]
        for start in range(0, len(records), 500):
            # Recheck IDs at commit time: another preview may have been confirmed first.
            portion = records[start:start+500]
            existing = set(str(x) for x in model.objects.filter(uid__in=[r['uid'] for r in portion if r.get('uid')]).values_list('uid', flat=True))
            model.objects.bulk_create([model(batch=batch, **record) for record in portion if record.get('uid') not in existing])
    batch.status, batch.imported_at = 'imported', timezone.now()
    if batch.format == 'legacy':
        batch.summary = dict(batch.summary, debt_policy=debt_policy)
    batch.save(update_fields=['status', 'imported_at', 'summary'])
    Activity.objects.create(title=gettext_noop('Excel импортирован'), detail=batch.filename, kind='import')
    return batch
