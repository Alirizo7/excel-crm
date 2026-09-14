"""Cell/formula bridge for the supported XLSX workbook format."""
import json
import math
import os
import re
import shutil
import subprocess
import uuid
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import to_excel, CALENDAR_WINDOWS_1900
from django.conf import settings
from django.core.exceptions import ValidationError
from django.utils.translation import gettext as _

MAX_ROWS, MAX_COLS, MAX_CELLS = 12000, 64, 200000
FORMULA_FUNCTIONS = {'SUM', 'MIN', 'MAX', 'AVERAGE', 'COUNT', 'COUNTA', 'IF', 'IFERROR', 'ROUND', 'ROUNDUP', 'ROUNDDOWN', 'ABS', 'NA'}
ERRORS = {'#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CYCLE!'}


def color(value, fallback=None):
    if value is not None and value.type == 'rgb' and isinstance(value.rgb, str):
        return '#' + value.rgb[-6:]
    if value is not None and value.type == 'indexed':
        from openpyxl.styles.colors import COLOR_INDEX
        if value.indexed < len(COLOR_INDEX):
            return '#' + COLOR_INDEX[value.indexed][-6:]
    return fallback


def read_xlsx(data, name):
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=False, keep_links=False)
    if workbook.epoch != CALENDAR_WINDOWS_1900:
        raise ValidationError(_('Книги с календарём 1904 пока не поддерживаются.'))
    if len(workbook.sheetnames) > 30:
        raise ValidationError(_('В книге слишком много листов.'))
    result = {'id': 'book-' + uuid.uuid4().hex, 'name': name, 'appVersion': '0.25.1',
              'locale': 'enUS', 'sheetOrder': [], 'sheets': {}, 'styles': {}}
    for i, ws in enumerate(workbook):
        if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS:
            raise ValidationError(_('Лист превышает допустимый размер.'))
        sid = f'sheet-{i}'
        sheet = {'id': sid, 'name': ws.title, 'rowCount': min(MAX_ROWS, max(ws.max_row + 100, 10100)),
                 'columnCount': max(ws.max_column, 12), 'cellData': {}, 'rowData': {}, 'columnData': {},
                 'defaultRowHeight': 26, 'defaultColumnWidth': 110, 'showGridlines': 1,
                 'rowHeader': {'width': 48}, 'columnHeader': {'height': 28}, 'hidden': int(ws.sheet_state != 'visible'),
                 'freeze': {'xSplit': 0, 'ySplit': 3, 'startRow': 3, 'startColumn': -1},
                 'mergeData': [{'startRow': m.min_row-1, 'endRow': m.max_row-1, 'startColumn': m.min_col-1, 'endColumn': m.max_col-1} for m in ws.merged_cells.ranges],
                 'custom': {'sourceName': ws.title, 'sourceRows': ws.max_row, 'sourceColumns': ws.max_column}}
        result['sheetOrder'].append(sid)
        for row in ws:
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                style_id = str(cell.style_id)
                if style_id not in result['styles']:
                    style = {'ff': cell.font.name or 'Arial', 'fs': cell.font.sz or 11,
                             'bl': int(bool(cell.font.b)), 'it': int(bool(cell.font.i)),
                             'cl': {'rgb': color(cell.font.color, '#18352a')},
                             'n': {'pattern': cell.number_format},
                             'ht': {'left': 1, 'center': 2, 'right': 3}.get(cell.alignment.horizontal, 0),
                             'vt': {'top': 1, 'center': 2, 'bottom': 3}.get(cell.alignment.vertical, 2),
                             'tb': 3 if cell.alignment.wrap_text else 1}
                    fill = color(cell.fill.fgColor) if cell.fill.patternType == 'solid' else None
                    if fill:
                        style['bg'] = {'rgb': fill}
                    borders = {}
                    for short, edge in [('t', 'top'), ('b', 'bottom'), ('l', 'left'), ('r', 'right')]:
                        side = getattr(cell.border, edge)
                        if side and side.style:
                            borders[short] = {'s': {'thin': 1, 'medium': 2, 'dashed': 3, 'dotted': 4, 'thick': 5, 'double': 6}.get(side.style, 1), 'cl': {'rgb': color(side.color, '#d3ded7')}}
                    if borders:
                        style['bd'] = borders
                    result['styles'][style_id] = style
                entry = {'s': style_id}
                if cell.data_type == 'f':
                    entry['f'] = re.sub(r'^=\+', '=', cell.value)
                elif cell.value is not None:
                    value = cell.value
                    if isinstance(value, (datetime, date)):
                        value = to_excel(value, workbook.epoch)
                    entry.update(v=value, t=3 if isinstance(value, bool) else 2 if isinstance(value, (int, float)) else 1)
                sheet['cellData'].setdefault(str(cell.row-1), {})[str(cell.column-1)] = entry
        for r in range(ws.max_row):
            dim = ws.row_dimensions.get(r+1)
            sheet['rowData'][str(r)] = {'custom': {'rowId': uuid.uuid4().hex},
                'h': max(24, (dim.height or 18)*4/3) if dim else 26, 'hd': int(bool(dim and dim.hidden))}
        for key, dim in ws.column_dimensions.items():
            col = column_index_from_string(key)-1
            for c in range(col, dim.max or col+1):
                if c < sheet['columnCount']:
                    sheet['columnData'][str(c)] = {'w': max(55, min(420, dim.width*7+5)), 'hd': int(dim.hidden)}
        result['sheets'][sid] = sheet
    return result


def node_binary():
    configured = os.getenv('WORKBOOK_NODE')
    if configured:
        return configured
    bundled = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
    return str(bundled) if bundled.exists() else (shutil.which('node') or 'node')


def calculate(data):
    try:
        bundled = settings.BASE_DIR/'scripts/calculate-workbook.bundle.mjs'
        script = bundled if bundled.is_file() else settings.BASE_DIR/'scripts/calculate-workbook.mjs'
        completed = subprocess.run([node_binary(), str(script)],
            input=json.dumps(data, ensure_ascii=False, allow_nan=False), text=True, capture_output=True,
            cwd=settings.BASE_DIR, timeout=30, check=True)
        return json.loads(completed.stdout)
    except (OSError, subprocess.SubprocessError, ValueError) as exc:
        raise ValidationError(_('Не удалось пересчитать книгу. Изменения не применены к учёту.')) from exc


def cells(data):
    for sid in data['sheetOrder']:
        sheet = data['sheets'][sid]
        for r, row in sheet.get('cellData', {}).items():
            for c, cell in row.items():
                if cell:
                    yield sid, int(r), int(c), cell


def formula_errors(data):
    return [{'sheet': data['sheets'][sid]['name'], 'cell': f'{get_column_letter(c+1)}{r+1}', 'error': entry['v'], 'legacy': bool(re.fullmatch(r"=(?:(?:'[^']+'|[\w ]+)!)?\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+", entry['f']))}
            for sid, r, c, entry in cells(data) if entry.get('f') and entry.get('v') in ERRORS]


def cell_content(cell):
    if cell.get('f'):
        return {'f': cell['f']}
    value = cell.get('v')
    if cell.get('p'):
        value = cell['p'].get('body', {}).get('dataStream', '').rstrip('\r\n')
    return {'v': value} if value is not None else {}


def safe_style(style):
    if not isinstance(style, dict):
        raise ValidationError(_('Некорректные стили книги.'))
    out = {}
    for key in ('ff',):
        if key in style:
            if not isinstance(style[key], str) or len(style[key]) > 100:
                raise ValidationError(_('Некорректные стили книги.'))
            out[key] = style[key]
    for key, maximum in [('fs', 100), ('bl', 1), ('it', 1), ('ht', 4), ('vt', 3), ('tb', 3)]:
        if style.get(key) is not None:
            value = style[key]
            if type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= maximum:
                raise ValidationError(_('Некорректные стили книги.'))
            out[key] = value
    for key in ('cl', 'bg'):
        part = style.get(key)
        if isinstance(part, dict) and isinstance(part.get('rgb'), str) and re.fullmatch(r'#[0-9a-fA-F]{6}', part['rgb']):
            out[key] = {'rgb': part['rgb']}
    pattern = style.get('n', {}).get('pattern') if isinstance(style.get('n'), dict) else None
    if pattern is not None:
        if not isinstance(pattern, str) or len(pattern) > 500:
            raise ValidationError(_('Некорректные стили книги.'))
        out['n'] = {'pattern': pattern}
    borders = style.get('bd')
    if isinstance(borders, dict):
        out['bd'] = {}
        for edge in ('t', 'b', 'l', 'r'):
            part = borders.get(edge)
            if isinstance(part, dict) and type(part.get('s')) == int and part['s'] in range(14):
                out['bd'][edge] = {'s': part['s'], **safe_style({'cl': part.get('cl')})}
    return out


def validate_snapshot(data, previous=None, allow_formulas=True):
    """Validate untrusted browser state before it reaches the calculation engine."""
    if not isinstance(data, dict) or not isinstance(data.get('sheets'), dict):
        raise ValidationError(_('Некорректные данные книги.'))
    order = data.get('sheetOrder')
    if not isinstance(order, list) or not 1 <= len(order) <= 30 or len(set(order)) != len(order) or set(order) != set(data['sheets']):
        raise ValidationError(_('Некорректный список листов.'))
    if previous and (data.get('id') != previous['id'] or set(order) != set(previous['sheetOrder'])):
        raise ValidationError(_('Состав листов исходной книги нельзя изменять.'))
    if previous and order != previous['sheetOrder']:
        raise ValidationError(_('Состав листов исходной книги нельзя изменять.'))
    if not isinstance(data.get('styles', {}), dict) or len(data.get('styles', {})) > 10000:
        raise ValidationError(_('Некорректные стили книги.'))
    data['styles'] = {key: safe_style(style) for key, style in data.get('styles', {}).items()}
    count = 0
    old_formulas = {}
    if previous and not allow_formulas:
        for sid, r, c, cell in cells(previous):
            rid = previous['sheets'][sid].get('rowData', {}).get(str(r), {}).get('custom', {}).get('rowId', f'index-{r}')
            if cell.get('f'):
                old_formulas[(sid, rid, c)] = cell['f']
    formulas = {}
    for sid in order:
        sheet = data['sheets'][sid]
        if not isinstance(sheet, dict) or sheet.get('id') != sid:
            raise ValidationError(_('Некорректный лист.'))
        if previous and sheet.get('name') != previous['sheets'][sid]['name']:
            raise ValidationError(_('Название исходного листа нельзя изменять.'))
        if not isinstance(sheet.get('name'), str) or not 1 <= len(sheet['name']) <= 31 or re.search(r'[\\/*?:\[\]]', sheet['name']):
            raise ValidationError(_('Некорректный лист.'))
        for key, maximum in [('rowCount', MAX_ROWS), ('columnCount', MAX_COLS)]:
            if type(sheet.get(key)) != int or not 1 <= sheet[key] <= maximum:
                raise ValidationError(_('Лист превышает допустимый размер.'))
        if previous and sheet['columnCount'] != previous['sheets'][sid]['columnCount']:
            raise ValidationError(_('Структуру колонок исходной книги нельзя изменять.'))
        for kind, dimension, key, fallback, maximum in [('rowData', 'rowCount', 'h', 26, 534), ('columnData', 'columnCount', 'w', 110, 1000)]:
            if not isinstance(sheet.get(kind, {}), dict):
                raise ValidationError(_('Некорректные строки листа.'))
            for index, dim in sheet.get(kind, {}).items():
                if not str(index).isdigit() or not 0 <= int(index) < sheet[dimension] or not isinstance(dim, dict):
                    raise ValidationError(_('Некорректные строки листа.'))
                value = dim.get(key, fallback)
                if type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= maximum:
                    raise ValidationError(_('Некорректные строки листа.'))
                if not isinstance(dim.get('custom', {}), dict):
                    raise ValidationError(_('Нарушена связь строк с учётом.'))
                sheet[kind][index] = {key: value, 'hd': int(bool(dim.get('hd'))), 'custom': {'rowId': dim.get('custom', {}).get('rowId')}}
        merges = sheet.get('mergeData', [])
        if not isinstance(merges, list) or len(merges) > 5000:
            raise ValidationError(_('Некорректный лист.'))
        for merge in merges:
            if not isinstance(merge, dict):
                raise ValidationError(_('Некорректный лист.'))
            for start, end, maximum in [('startRow', 'endRow', sheet['rowCount']), ('startColumn', 'endColumn', sheet['columnCount'])]:
                if type(merge.get(start)) != int or type(merge.get(end)) != int or not 0 <= merge[start] <= merge[end] < maximum:
                    raise ValidationError(_('Некорректный лист.'))
        freeze = sheet.get('freeze') or {}
        if not isinstance(freeze, dict):
            raise ValidationError(_('Некорректный лист.'))
        for key in ('xSplit', 'ySplit', 'startRow', 'startColumn'):
            if key in freeze and (type(freeze[key]) != int or not -1 <= freeze[key] <= MAX_ROWS):
                raise ValidationError(_('Некорректный лист.'))
        if not isinstance(sheet.get('cellData', {}), dict):
            raise ValidationError(_('Некорректное содержимое ячейки.'))
        row_ids = set()
        for r, dim in sheet.get('rowData', {}).items():
            if not str(r).isdigit() or not 0 <= int(r) < sheet['rowCount'] or not isinstance(dim, dict):
                raise ValidationError(_('Некорректные строки листа.'))
            rid = dim.get('custom', {}).get('rowId')
            if rid:
                if not isinstance(rid, str) or not re.fullmatch('[a-f0-9]{32}', rid) or rid in row_ids:
                    raise ValidationError(_('Нарушена связь строк с учётом.'))
                row_ids.add(rid)
        if previous and not allow_formulas:
            old_ids = {r: d.get('custom', {}).get('rowId') for r, d in previous['sheets'][sid].get('rowData', {}).items()}
            new_ids = {r: d.get('custom', {}).get('rowId') for r, d in sheet.get('rowData', {}).items()}
            if any(new_ids.get(r) != rid for r, rid in old_ids.items() if rid):
                raise ValidationError(_('Изменять структуру может только ответственный за формулы.'))
        for r, row in sheet.get('cellData', {}).items():
            if not str(r).isdigit() or not 0 <= int(r) < sheet['rowCount'] or not isinstance(row, dict):
                raise ValidationError(_('Некорректный адрес ячейки.'))
            for c, cell in list(row.items()):
                if not str(c).isdigit() or not 0 <= int(c) < sheet['columnCount']:
                    raise ValidationError(_('Некорректный адрес ячейки.'))
                if cell is None:
                    del row[c]
                    continue
                if not isinstance(cell, dict):
                    raise ValidationError(_('Некорректное содержимое ячейки.'))
                count += 1
                if count > MAX_CELLS:
                    raise ValidationError(_('В книге слишком много ячеек.'))
                content = cell_content(cell)
                f = content.get('f')
                if f:
                    if not isinstance(f, str) or len(f) > 2000 or not f.startswith('=') or '[' in f or '\\' in f:
                        raise ValidationError(_('Формула слишком длинная или содержит внешнюю ссылку.'))
                    functions = {x.upper() for x in re.findall(r'([A-Za-z_][A-Za-z0-9_.]*)\s*\(', f)}
                    if functions - FORMULA_FUNCTIONS:
                        raise ValidationError(_('Формула содержит неподдерживаемую функцию: ') + ', '.join(sorted(functions - FORMULA_FUNCTIONS)))
                    for col, nr in re.findall(r'\$?([A-Z]{1,3})\$?(\d+)', f):
                        if int(nr) > MAX_ROWS or column_index_from_string(col) > MAX_COLS:
                            raise ValidationError(_('Ссылка в формуле превышает допустимый размер книги.'))
                    rid = sheet.get('rowData', {}).get(r, {}).get('custom', {}).get('rowId', f'index-{r}')
                    formulas[(sid, rid, int(c))] = f
                else:
                    v = content.get('v')
                    if v is not None and (type(v) not in (str, int, float, bool) or isinstance(v, str) and len(v) > 10000 or isinstance(v, (float, int)) and not isinstance(v, bool) and (not math.isfinite(v) or abs(v) > 1e20)):
                        raise ValidationError(_('Недопустимое значение ячейки.'))
                # Rich document runs, URLs and client-calculated formula values are not authoritative.
                row[c] = {**content, **({'s': cell['s']} if cell.get('s') is not None else {})}
                if isinstance(row[c].get('s'), dict):
                    row[c]['s'] = safe_style(row[c]['s'])
                elif row[c].get('s') is not None and (not isinstance(row[c]['s'], str) or row[c]['s'] not in data['styles']):
                    row[c].pop('s')
                if not f and content.get('v') is not None:
                    row[c]['t'] = 3 if isinstance(content['v'], bool) else 2 if isinstance(content['v'], (int, float)) else 1
        for r in sheet.get('cellData', {}):
            custom = sheet.setdefault('rowData', {}).setdefault(r, {}).setdefault('custom', {})
            if not custom.get('rowId'):
                custom['rowId'] = uuid.uuid4().hex
    if previous and not allow_formulas and old_formulas != formulas:
        raise ValidationError(_('Формулы защищены. Изменять их может только ответственный.'))
    if not isinstance(data.get('styles', {}), dict) or len(data.get('styles', {})) > 10000:
        raise ValidationError(_('Некорректные стили книги.'))
    # No plugins or embedded external resources are accepted from the client.
    data.pop('resources', None)
    if not any(not sheet.get('hidden') for sheet in data['sheets'].values()):
        raise ValidationError(_('В книге должен оставаться видимый лист.'))
    return data


def write_xlsx(data, cached=True):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    style_cache = {}
    for sid in data['sheetOrder']:
        sheet = data['sheets'][sid]
        ws = workbook.create_sheet(sheet['name'])
        for r, row in sheet.get('cellData', {}).items():
            for c, entry in row.items():
                if not entry:
                    continue
                cell = ws.cell(int(r)+1, int(c)+1)
                cell.value = entry.get('f') or entry.get('v')
                if not entry.get('f') and isinstance(entry.get('v'), str):
                    cell.data_type = 's'
                reference = entry.get('s')
                style = data.get('styles', {}).get(reference, {}) if isinstance(reference, str) else (reference or {})
                if isinstance(style, dict):
                    key = ('id', reference) if isinstance(reference, str) else ('inline', json.dumps(style, sort_keys=True))
                    if key not in style_cache:
                        def rgb(part, fallback='000000'):
                            value = style.get(part, {})
                            value = value.get('rgb', '') if isinstance(value, dict) else ''
                            return value.lstrip('#') if re.fullmatch(r'#[0-9A-Fa-f]{6}', value) else fallback
                        parts = {
                            'font': Font(name=style.get('ff', 'Arial'), size=style.get('fs', 11), bold=bool(style.get('bl')), italic=bool(style.get('it')), color=rgb('cl')),
                            'fill': PatternFill('solid', fgColor=rgb('bg', 'FFFFFF')) if style.get('bg') else PatternFill(),
                            'alignment': Alignment(horizontal={1: 'left', 2: 'center', 3: 'right'}.get(style.get('ht')), vertical={1: 'top', 2: 'center', 3: 'bottom'}.get(style.get('vt'), 'center'), wrap_text=style.get('tb') == 3),
                        }
                        for attr, value in parts.items():
                            setattr(cell, attr, value)
                        cell.number_format = style.get('n', {}).get('pattern', 'General') if isinstance(style.get('n'), dict) else 'General'
                        sides = {}
                        for short, edge in [('t', 'top'), ('b', 'bottom'), ('l', 'left'), ('r', 'right')]:
                            b = (style.get('bd') or {}).get(short)
                            if isinstance(b, dict):
                                co = (b.get('cl') or {}).get('rgb', '#D3DED7')
                                co = co.lstrip('#') if re.fullmatch(r'#[0-9A-Fa-f]{6}', co) else 'D3DED7'
                                sides[edge] = Side(style={1:'thin', 2:'medium', 3:'dashed', 4:'dotted', 5:'thick', 6:'double'}.get(b.get('s'), 'thin'), color=co)
                        if sides:
                            cell.border = Border(**sides)
                        # Reuse registered style IDs instead of hashing every font and
                        # border for every cell. Copy the array: merging may modify it.
                        style_cache[key] = copy(cell._style)
                    else:
                        cell._style = copy(style_cache[key])
        for r, dim in sheet.get('rowData', {}).items():
            # Empty row IDs should not enlarge the exported used range.
            if str(r) in sheet.get('cellData', {}) or dim.get('hd'):
                ws.row_dimensions[int(r)+1].height = min(400, max(12, dim.get('h', 26)*.75))
                ws.row_dimensions[int(r)+1].hidden = bool(dim.get('hd'))
        for c, dim in sheet.get('columnData', {}).items():
            ws.column_dimensions[get_column_letter(int(c)+1)].width = max(4, (dim.get('w', 110)-5)/7)
            ws.column_dimensions[get_column_letter(int(c)+1)].hidden = bool(dim.get('hd'))
        for merge in sheet.get('mergeData', []):
            ws.merge_cells(start_row=merge['startRow']+1, end_row=merge['endRow']+1, start_column=merge['startColumn']+1, end_column=merge['endColumn']+1)
        freeze = sheet.get('freeze') or {}
        if freeze.get('ySplit', 0) or freeze.get('xSplit', 0):
            ws.freeze_panes = f"{get_column_letter(max(0, freeze.get('xSplit', 0))+1)}{max(0, freeze.get('ySplit', 0))+1}"
        ws.sheet_state = 'hidden' if sheet.get('hidden') else 'visible'
    output = BytesIO()
    workbook.save(output)
    if not cached:
        return output.getvalue()
    # openpyxl intentionally does not write formula caches. Store verified engine
    # results in OOXML so Excel, previews and the existing importer see the same values.
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    output2 = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(output2, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            match = re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml', item.filename)
            if match:
                sid = data['sheetOrder'][int(match.group(1))-1]
                sheet = data['sheets'][sid]
                root = ET.fromstring(payload)
                for cell in root.iter(f'{{{ns}}}c'):
                    if cell.find(f'{{{ns}}}f') is None:
                        continue
                    col, row = re.fullmatch(r'([A-Z]+)(\d+)', cell.attrib['r']).groups()
                    value = sheet.get('cellData', {}).get(str(int(row)-1), {}).get(str(column_index_from_string(col)-1), {}).get('v')
                    v = cell.find(f'{{{ns}}}v')
                    if v is None:
                        v = ET.SubElement(cell, f'{{{ns}}}v')
                    cell.attrib.pop('t', None)
                    if isinstance(value, str):
                        cell.set('t', 'e' if value in ERRORS else 'str')
                        v.text = value
                    elif isinstance(value, bool):
                        cell.set('t', 'b'); v.text = '1' if value else '0'
                    elif value is not None:
                        v.text = str(value)
                payload = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(item, payload)
    return output2.getvalue()
