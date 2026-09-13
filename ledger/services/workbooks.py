import copy
import hashlib
import json
import uuid
import zlib
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.utils import timezone
from django.utils.translation import gettext as _, gettext_noop

from ledger.models import (WorkingWorkbook, WorkbookVersion, ImportBatch, Operation, Delivery,
                           PartnerBalance, RecordChange, Activity)
from .changes import workspace_transaction, snapshot, changed, log_change, partner_for
from .importing import inspect_workbook, num
from .workbook_files import read_xlsx, validate_snapshot, calculate, write_xlsx, cells, cell_content, formula_errors

MODELS = {'operations': Operation, 'deliveries': Delivery, 'balances': PartnerBalance}
FIELDS = {
    'operations': ['date', 'kind', 'amount', 'description', 'category', 'partner'],
    'deliveries': ['date', 'direction', 'partner', 'vehicle', 'gross', 'tare', 'discount', 'price', 'notes'],
    'balances': ['name', 'balance', 'note'],
}
NUMBERS = {'amount', 'gross', 'tare', 'discount', 'price', 'balance'}


def error_keys(data):
    from .workbook_files import ERRORS
    return [f"{sid}:{data['sheets'][sid].get('rowData', {}).get(str(r), {}).get('custom', {}).get('rowId', r)}:{c}:{cell['v']}"
            for sid, r, c, cell in cells(data) if cell.get('f') and cell.get('v') in ERRORS]


def numeric_inputs(data):
    columns = {'Приход-Расход': (1, 4), 'тура ака': (5, 6, 8, 10),
               'дафтарча': (5, 6, 8, 9), 'Чугун': (7, 8, 10, 12), 'Кунлик': (3, 4)}
    result = {}
    for sid, r, c, cell in cells(data):
        sheet = data['sheets'][sid]
        if r < 3 or c not in columns.get(sheet['name'], ()):
            continue
        rid = sheet.get('rowData', {}).get(str(r), {}).get('custom', {}).get('rowId', r)
        value = cell.get('v')
        if value is not None and value != '':
            result[f'{sid}:{rid}:{c}'] = {'value': value, 'sheet': sheet['name'], 'row': r+1, 'column': c+1}
    return result


class WorkbookConflict(ValidationError):
    pass


def current_book(book):
    if book.batch.status != ImportBatch.Status.IMPORTED:
        raise ValidationError(_('Эта книга архивирована новым импортом. Доступны просмотр и скачивание.'))


def canonical(kind, values):
    result = {}
    for field in FIELDS[kind]:
        value = values.get(field)
        if field in NUMBERS and value is not None:
            value = format(Decimal(str(value)).normalize(), 'f')
        elif field == 'date':
            value = str(value) if value else None
        result[field] = value
    return result


def project(data, batch):
    parsed = inspect_workbook(write_xlsx(data), batch.filename, batch.report_date, skip_existing_ids=False)
    names = {s['name']: (sid, s) for sid, s in data['sheets'].items()}
    result = {}
    for kind in MODELS:
        for record in parsed[kind]:
            sid, sheet = names[record['source_sheet']]
            r = str(record['source_row']-1)
            rid = sheet['rowData'].get(r, {}).get('custom', {}).get('rowId')
            if not rid:
                raise ValidationError(_('Строка потеряла связь с учётом. Обновите страницу.'))
            discriminator = record['kind'] if kind == 'operations' else ''
            key = f'{kind}:{sid}:{rid}:{discriminator}'
            result[key] = {'kind': kind, 'sheet': record['source_sheet'], 'row': record['source_row'],
                           'values': canonical(kind, record), 'uid': record.get('uid')}
            if kind == 'deliveries' and parsed['format'] == 'legacy':
                # Keep the original formula outputs in the projection too. CRM
                # delivery totals follow its weight/discount/price contract.
                columns = {'тура ака': (7, 9, 11), 'дафтарча': (7, None, 10), 'Чугун': (9, 11, 13)}
                if entry_columns := columns.get(sheet['name']):
                    result[key]['outputs'] = [cell_content(sheet.get('cellData', {}).get(r, {}).get(str(c), {})) for c in entry_columns]
                    result[key]['output_values'] = [str(num(sheet.get('cellData', {}).get(r, {}).get(str(c), {}).get('v'))) for c in entry_columns]
    return result, parsed['issues']


def version(book, before, actor, reason):
    old = {(sid, r, c): cell_content(cell) for sid, r, c, cell in cells(before)} if before else {}
    new = {(sid, r, c): cell_content(cell) for sid, r, c, cell in cells(book.data)}
    from openpyxl.utils import get_column_letter
    deltas = [{'sheet': book.data['sheets'][sid]['name'], 'cell': f'{get_column_letter(c+1)}{r+1}',
               'before': old.get((sid, r, c), {}), 'after': new.get((sid, r, c), {})}
              for sid, r, c in sorted(set(old) | set(new)) if old.get((sid, r, c), {}) != new.get((sid, r, c), {})] if before else []
    WorkbookVersion.objects.create(workbook=book, revision=book.revision, actor=actor, reason=reason,
        payload=zlib.compress(json.dumps(book.data, ensure_ascii=False).encode(), level=6), changes=deltas)


def create_book(batch_id, actor):
    batch = ImportBatch.objects.get(pk=batch_id, status=ImportBatch.Status.IMPORTED)
    existing = WorkingWorkbook.objects.filter(batch=batch).first()
    if existing:
        return existing
    with batch.file.open('rb') as file:
        data = validate_snapshot(read_xlsx(file.read(), batch.filename))
    data = calculate(data)
    projection, issues = project(data, batch)
    with workspace_transaction():
        batch.refresh_from_db()
        if batch.status != ImportBatch.Status.IMPORTED:
            raise WorkbookConflict(_('Импорт уже заменён. Откройте актуальную книгу.'))
        existing = WorkingWorkbook.objects.filter(batch=batch).first()
        if existing:
            return existing
        baseline = {}
        records = {kind: list(model.objects.filter(batch=batch)) for kind, model in MODELS.items()}
        # Use original operation sides when a CRM form has already changed its type.
        old_kinds = {}
        for change in RecordChange.objects.filter(kind='operations', record_id__in=[r.pk for r in records['operations']]).order_by('created_at'):
            if change.before.get('kind'):
                old_kinds.setdefault(change.record_id, change.before['kind'])
        for key, entry in projection.items():
            kind = entry['kind']
            matches = [r for r in records[kind] if r.source_sheet == entry['sheet'] and r.source_row == entry['row']
                       and (kind != 'operations' or old_kinds.get(r.pk, r.kind) == entry['values']['kind'])]
            if entry.get('uid'):
                matches = [r for r in records[kind] if str(r.uid) == entry['uid']]
            if len(matches) == 1:
                baseline[key] = {'uid': str(matches[0].uid), 'values': canonical(kind, snapshot(matches[0])), 'active': matches[0].active}
        book = WorkingWorkbook.objects.create(batch=batch, data=data, projection={'records': projection, 'issues': issues, 'formula_errors': error_keys(data), 'numeric_inputs': numeric_inputs(data)}, ledger_baseline=baseline)
        version(book, None, actor, 'create')
        Activity.objects.create(title=gettext_noop('Рабочая книга создана'), detail=batch.filename, kind='workbook')
        return book


def save_book(book_id, data, revision, actor, allow_formulas, reason='edit'):
    book = WorkingWorkbook.objects.select_related('batch').get(pk=book_id)
    current_book(book)
    if book.revision != revision:
        raise WorkbookConflict(_('Книгу уже изменили в другой вкладке. Ваши правки сохранены на экране; обновите книгу после сверки.'))
    data = validate_snapshot(copy.deepcopy(data), book.data, allow_formulas)
    data = calculate(data)
    with workspace_transaction():
        book = WorkingWorkbook.objects.select_related('batch').get(pk=book_id)
        current_book(book)
        if book.revision != revision:
            raise WorkbookConflict(_('Версия книги изменилась во время сохранения. Обновите страницу после сверки правок.'))
        before = book.data
        if data == before:
            return book
        book.data = data
        book.revision += 1
        book.save(update_fields=['data', 'revision', 'updated_at'])
        version(book, before, actor, reason)
        return book


def restore_book(book_id, target, revision, actor, allow_formulas):
    item = WorkbookVersion.objects.get(workbook_id=book_id, revision=target)
    data = json.loads(zlib.decompress(item.payload))
    return save_book(book_id, data, revision, actor, allow_formulas, reason='restore')


def _plan(book, projected, issues):
    old = {key: {**entry, 'values': canonical(entry['kind'], entry['values'])} for key, entry in book.projection.get('records', {}).items()}
    baseline = book.ledger_baseline
    lookup = {kind: {str(r.uid): r for r in model.objects.filter(uid__in=[b['uid'] for key, b in baseline.items() if key.startswith(kind+':')])} for kind, model in MODELS.items()}
    existing_issues = {json.dumps(i, sort_keys=True) for i in book.projection.get('issues', [])}
    blocked = [i['message'] for i in issues if json.dumps(i, sort_keys=True) not in existing_issues
               and (i['level'] == 'error' or 'Поставка пропущена' in i['message'])]
    if set(error_keys(book.data)) - set(book.projection.get('formula_errors', [])):
        blocked.append(_('В книге появились новые ошибки формул. Исправьте их перед применением.'))
    before_inputs = book.projection.get('numeric_inputs', {})
    if book.batch.format == 'legacy':
        from openpyxl.utils import get_column_letter
        for key, entry in numeric_inputs(book.data).items():
            if entry['value'] != before_inputs.get(key, {}).get('value') and num(entry['value']) is None:
                blocked.append(f"{entry['sheet']}!{get_column_letter(entry['column'])}{entry['row']}: " + _('В числовом поле должно быть число. Для удаления значения очистите ячейку.'))
    plans = []
    for key in sorted(set(old) | set(projected)):
        prev, now = old.get(key), projected.get(key)
        if prev and now and prev['values'] == now['values'] and prev.get('output_values') == now.get('output_values'):
            continue
        entry = now or prev
        kind = entry['kind']
        base = baseline.get(key)
        record = lookup[kind].get(base['uid']) if base else None
        fields = [f for f in FIELDS[kind] if not prev or not now or prev['values'].get(f) != now['values'].get(f)]
        errors = []
        if base and record is None:
            errors.append(_('Связанная запись больше недоступна.'))
        if record:
            current = canonical(kind, snapshot(record))
            compare_fields = FIELDS[kind] if not now else fields
            if record.archived_at or record.active != base['active'] or any(current[f] != canonical(kind, base['values'])[f] for f in compare_fields):
                errors.append(_('Эта запись изменена в CRM. Сверьте правки в карточке записи.'))
        if not now and not record:
            continue
        before = snapshot(record) if record else {}
        candidate = copy.deepcopy(record) if record else MODELS[kind](batch=book.batch, source_sheet=entry['sheet'], source_row=entry['row'])
        if now:
            candidate.active = True
            candidate.deleted_at = None
            for field in fields:
                value = now['values'][field]
                if field in NUMBERS:
                    value = Decimal(value)
                elif field == 'date' and value:
                    from datetime import date
                    value = date.fromisoformat(value)
                setattr(candidate, field, value)
            if kind == 'deliveries':
                candidate.net, candidate.clean_weight, candidate.amount = Delivery.calculate(candidate.gross, candidate.tare, candidate.discount, candidate.price)
                if now.get('outputs'):
                    expected_outputs = (candidate.net, candidate.clean_weight, candidate.amount)
                    for value, expected in zip(now['output_values'], expected_outputs):
                        if value != 'None' and abs(Decimal(value) - expected) > Decimal('.01'):
                            errors.append(_('Формула поставки расходится с расчётом веса, скидки и цены в CRM. Сверьте формулу и исходные поля.'))
                            break
            if kind == 'balances' and candidate.paid_amount:
                if abs(candidate.balance) < candidate.paid_amount or candidate.balance * record.balance <= 0 or candidate.name != record.name:
                    errors.append(_('Есть погашения: нельзя менять сторону, контрагента или уменьшать долг ниже оплаченной суммы.'))
            try:
                candidate.full_clean(exclude=['uid', 'counterparty'], validate_unique=False)
            except ValidationError as exc:
                errors.extend(exc.messages)
        elif kind == 'balances' and record.paid_amount:
            errors.append(_('Сначала отмените погашения этого долга.'))
        plans.append({'key': key, 'kind': kind, 'sheet': entry['sheet'], 'row': entry['row'],
            'action': 'delete' if not now else ('restore' if not record.active else 'edit') if record else 'create',
            'fields': fields, 'record': record, 'candidate': candidate, 'before': before,
            'after': canonical(kind, snapshot(candidate)) if now else {}, 'errors': errors})
    fingerprint = [{'key': p['key'], 'action': p['action'], 'before': p['before'], 'after': p['after'], 'errors': p['errors']} for p in plans]
    token = hashlib.sha256(json.dumps([book.revision, fingerprint, blocked], sort_keys=True, default=str).encode()).hexdigest()
    return plans, blocked, token


def preview_book(book):
    current_book(book)
    projected, issues = project(book.data, book.batch)
    plans, blocked, token = _plan(book, projected, issues)
    return plans, blocked, token, projected, issues


def apply_book(book_id, revision, token, actor, *, prepared=None):
    # The one-click view has already projected this exact workbook revision.
    # Reuse that work, but always recheck the revision and CRM state under lock.
    if prepared is None:
        book = WorkingWorkbook.objects.select_related('batch').get(pk=book_id)
        if book.revision != revision:
            raise WorkbookConflict(_('Книга изменилась. Заново откройте проверку изменений.'))
        projected, issues = project(book.data, book.batch)
    else:
        prepared_revision, projected, issues = prepared
        if prepared_revision != revision:
            raise WorkbookConflict(_('Книга изменилась. Заново откройте проверку изменений.'))
    with workspace_transaction():
        book = WorkingWorkbook.objects.select_related('batch').get(pk=book_id)
        current_book(book)
        if revision != book.revision:
            raise WorkbookConflict(_('Книга изменилась. Заново откройте проверку изменений.'))
        plans, blocked, expected = _plan(book, projected, issues)
        if token != expected:
            raise WorkbookConflict(_('Данные учёта изменились. Заново проверьте изменения перед применением.'))
        if blocked or any(p['errors'] for p in plans):
            raise ValidationError(_('Исправьте отмеченные ошибки перед применением к учёту.'))
        for plan in plans:
            obj = plan['candidate']
            if plan['action'] == 'delete':
                obj.active = False
                obj.deleted_at = timezone.now()
            elif plan['kind'] == 'balances':
                obj.counterparty = partner_for(obj.name)
            changed(obj)
            obj.save()
            log_change(obj, 'partners' if plan['kind'] == 'balances' else plan['kind'], plan['before'], plan['action'], actor)
            book.ledger_baseline[plan['key']] = {'uid': str(obj.uid), 'values': canonical(plan['kind'], snapshot(obj)), 'active': obj.active}
        book.projection = {'records': projected, 'issues': issues, 'formula_errors': error_keys(book.data), 'numeric_inputs': numeric_inputs(book.data)}
        book.applied_revision = book.revision
        book.save(update_fields=['projection', 'ledger_baseline', 'applied_revision', 'updated_at'])
        Activity.objects.create(title=gettext_noop('Рабочая книга применена к учёту'), detail=book.batch.filename, kind='workbook')
        return len(plans)
