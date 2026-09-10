import uuid
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone, translation
from openpyxl import load_workbook

from .models import Counterparty, DebtPayment, PartnerBalance, Operation, Delivery, RecordChange, ImportBatch
from .forms import DebtForm
from .services.changes import partner_for
from .services.debts import debt_rows, debt_stats, record_payment, cancel_payment
from .services.importing import stage_import, commit_import
from .tests import legacy_bytes, book_bytes


class DebtWorkflowTests(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.get(username='admin'))
        self.media = TemporaryDirectory()
        self.override = override_settings(MEDIA_ROOT=self.media.name)
        self.override.enable()
        self.addCleanup(self.media.cleanup)
        self.addCleanup(self.override.disable)
        self.partner = partner_for('Тестовый партнёр')
        self.debt = PartnerBalance.objects.create(name=self.partner.name, counterparty=self.partner, balance=Decimal('1000.50'))

    def pay(self, amount='200.25', debt=None, **extra):
        debt = debt or self.debt
        debt.refresh_from_db()
        data = dict(amount=Decimal(amount), date=timezone.localdate(), note='Тест оплаты',
                    request_key=uuid.uuid4(), revision=debt.revision, actor='tester')
        data.update(extra)
        return record_payment(debt.pk, **data)

    def test_partial_full_and_both_sides_statistics(self):
        payable = PartnerBalance.objects.create(name=self.partner.name, counterparty=self.partner, balance=Decimal('-500.75'))
        self.pay()
        self.pay('100.25', debt=payable)
        self.debt.refresh_from_db(); payable.refresh_from_db()
        self.assertEqual(self.debt.outstanding, Decimal('800.25'))
        self.assertEqual(payable.remaining, Decimal('-400.50'))
        stats = debt_stats(debt_rows().filter(active=True))
        self.assertEqual(stats['receivable'], Decimal('800.25'))
        self.assertEqual(stats['payable'], Decimal('400.50'))
        self.assertEqual(stats['net'], Decimal('399.75'))
        self.pay('800.25')
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.outstanding, 0)
        self.assertEqual(Operation.objects.count(), 0)
        self.assertContains(self.client.get(reverse('debt_detail', args=[self.debt.pk])), 'Погашен')
        self.assertEqual(debt_stats(debt_rows().filter(active=True))['closed_count'], 1)

    def test_duplicate_http_submission_is_idempotent_even_after_full_payment(self):
        values = {'date': str(timezone.localdate()), 'amount': '1000.50', 'note': 'Полностью',
                  'request_key': str(uuid.uuid4()), 'revision': 0}
        url = reverse('payment_new', args=[self.debt.pk])
        for _ in range(2):
            self.assertRedirects(self.client.post(url, values), reverse('debt_detail', args=[self.debt.pk]))
        self.assertEqual(DebtPayment.objects.count(), 1)
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.paid_amount, Decimal('1000.50'))
        values['amount'] = '10'
        self.assertContains(self.client.post(url, values), 'Этот запрос уже использован')
        self.assertEqual(DebtPayment.objects.count(), 1)

    def test_overpayment_zero_future_date_and_stale_payment_are_rejected(self):
        for amount in ['0', '-1', '1000.51']:
            with self.subTest(amount=amount), self.assertRaises(ValidationError):
                self.pay(amount)
        with self.assertRaises(ValidationError):
            self.pay(date=timezone.localdate() + timedelta(days=1))
        self.pay('100')
        with self.assertRaises(ValidationError):
            self.pay('50', revision=0)
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.paid_amount, 100)
        self.assertEqual(DebtPayment.objects.count(), 1)

    def test_cancel_restores_balance_and_retains_history(self):
        payment, _ = self.pay()
        self.debt.refresh_from_db()
        with self.assertRaises(ValidationError):
            cancel_payment(payment.pk, reason='', revision=self.debt.revision, actor='tester')
        self.assertTrue(cancel_payment(payment.pk, reason='Ошибочная сумма', revision=self.debt.revision, actor='tester'))
        self.assertFalse(cancel_payment(payment.pk, reason='Повтор', revision=self.debt.revision, actor='tester'))
        self.debt.refresh_from_db(); payment.refresh_from_db()
        self.assertEqual(self.debt.paid_amount, 0)
        self.assertEqual(payment.cancel_reason, 'Ошибочная сумма')
        self.assertEqual(RecordChange.objects.filter(kind='partners').count(), 2)
        response = self.client.get(reverse('partners'))
        self.assertEqual(response.context['stats']['paid'], 0)
        self.assertContains(response, 'Отменена')

    def test_edit_debt_preserves_payments_and_rejects_invalid_changes(self):
        self.pay('300')
        self.debt.refresh_from_db()
        values = dict(name=self.partner.name, side='receivable', principal='1300.50', note='За металл',
                      due_date=str(timezone.localdate()-timedelta(days=1)), revision=self.debt.revision)
        self.assertEqual(self.client.post(reverse('debt_edit', args=[self.debt.pk]), values).status_code, 302)
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.outstanding, Decimal('1000.50'))
        self.assertTrue(self.debt.overdue)
        self.assertContains(self.client.post(reverse('debt_edit', args=[self.debt.pk]), values), 'Запись уже изменена')
        for fields in [dict(principal='299'), dict(side='payable'), dict(name='Другой человек')]:
            response = self.client.post(reverse('debt_edit', args=[self.debt.pk]), dict(values, **fields))
            self.assertEqual(response.status_code, 200)
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.name, self.partner.name)
        self.assertEqual(self.debt.balance, Decimal('1300.50'))
        response = self.client.post(reverse('debt_delete', args=[self.debt.pk]), {'revision': self.debt.revision}, follow=True)
        self.assertContains(response, 'Сначала отмените погашения')
        self.debt.refresh_from_db(); self.assertTrue(self.debt.active)

    def test_create_delete_restore_debt_and_counterparty_grouping(self):
        response = self.client.post(reverse('debt_new'), {'name': '  ТЕСТОВЫЙ   ПАРТНЁР ', 'side': 'payable', 'principal': '42.50', 'note': 'Тест'})
        self.assertEqual(response.status_code, 302)
        debt = PartnerBalance.objects.get(balance=Decimal('-42.50'))
        self.assertEqual(debt.counterparty_id, self.partner.pk)
        for url in ['debt_delete', 'debt_restore']:
            self.assertEqual(self.client.post(reverse(url, args=[debt.pk]), {'revision': debt.revision}).status_code, 302)
            debt.refresh_from_db()
        self.assertTrue(debt.active)
        self.assertEqual(Counterparty.objects.count(), 1)
        self.assertEqual(self.client.get(reverse('partner_detail', args=[self.partner.pk])).context['debts'].count(), 2)

    def test_export_current_balances_and_cancelled_payments_in_both_languages(self):
        payment, _ = self.pay('250.25')
        for lang in ['ru', 'uz']:
            self.client.cookies['django_language'] = lang
            response = self.client.get(reverse('export_download', args=['partners']))
            book = load_workbook(BytesIO(response.content))
            self.assertEqual(book['Взаиморасчёты']['D2'].value, 750.25)
            self.assertEqual(book['Взаиморасчёты']['G2'].value, 250.25)
            self.assertEqual(book['Погашения долгов']['C2'].value, 250.25)
        self.debt.refresh_from_db()
        cancel_payment(payment.pk, reason='Ошибка', revision=self.debt.revision, actor='tester')
        book = load_workbook(BytesIO(self.client.get(reverse('export_download', args=['report'])).content))
        self.assertEqual(book['Погашения долгов']['F2'].value, 'Отменена')
        self.assertEqual(book['Взаиморасчёты']['D2'].value, 1000.5)

    def test_reimport_requires_choice_and_keeps_or_archives_paid_debts(self):
        first, _ = stage_import(legacy_bytes(), 'first.xlsx'); commit_import(first.pk)
        debt = PartnerBalance.objects.get(batch=first, name='Партнёр А')
        payment, _ = self.pay('200', debt=debt)
        second, _ = stage_import(legacy_bytes(income=700), 'second.xlsx')
        with self.assertRaises(ValidationError): commit_import(second.pk)
        commit_import(second.pk, debt_policy='keep')
        debt.refresh_from_db()
        self.assertTrue(debt.active); self.assertIsNone(debt.archived_at)
        self.assertEqual(debt.outstanding, 550)
        self.assertFalse(PartnerBalance.objects.filter(batch=second).exists())
        book = load_workbook(BytesIO(legacy_bytes(income=800)))
        book['Кунлик']['D5'] = 800
        third, _ = stage_import(book_bytes(book), 'third.xlsx')
        with self.assertRaises(ValidationError): commit_import(third.pk, debt_policy='replace')
        commit_import(third.pk, debt_policy='replace', replace_local_changes=True)
        debt.refresh_from_db()
        self.assertFalse(debt.active); self.assertIsNotNone(debt.archived_at)
        current = PartnerBalance.objects.get(batch=third, name='Партнёр А')
        self.assertEqual(current.outstanding, 550)
        self.assertEqual(current.counterparty_id, debt.counterparty_id)
        self.assertEqual(DebtPayment.objects.get().pk, payment.pk)
        self.assertContains(self.client.get(reverse('partner_detail', args=[current.counterparty_id])), '200,00')
        with self.assertRaises(ValidationError):
            cancel_payment(payment.pk, reason='Архив', revision=debt.revision, actor='tester')

    def test_imported_edits_audit_delete_restore_and_new_snapshot(self):
        batch, _ = stage_import(legacy_bytes(), 'first.xlsx'); commit_import(batch.pk)
        item = Operation.objects.filter(batch=batch, kind='income').first()
        before_payload = batch.payload
        uid, original_date = item.uid, item.date
        response = self.client.get(reverse('record_edit', args=['operations', item.pk]))
        self.assertEqual(response.context['form'].instance.date, original_date)
        values = {'date': str(original_date), 'kind': item.kind, 'amount': '123.45', 'description': 'Исправлено', 'category': 'Прочее', 'partner': '', 'revision': 0}
        self.assertEqual(self.client.post(reverse('record_edit', args=['operations', item.pk]), values).status_code, 302)
        item.refresh_from_db(); batch.refresh_from_db()
        self.assertEqual(item.amount, Decimal('123.45')); self.assertEqual(item.uid, uid)
        self.assertEqual(batch.payload, before_payload)
        change = RecordChange.objects.get(kind='operations', record_id=item.pk)
        self.assertEqual(change.after['description'], 'Исправлено')
        self.assertContains(self.client.post(reverse('record_edit', args=['operations', item.pk]), values), 'Запись уже изменена')
        self.client.post(reverse('record_delete', args=['operations', item.pk]), {'revision': item.revision})
        item.refresh_from_db(); self.assertFalse(item.active)
        exported = load_workbook(BytesIO(self.client.get(reverse('export_download', args=['operations']), {'status': 'deleted'}).content))
        self.assertNotIn(str(uid), [row[0] for row in exported.active.iter_rows(min_row=2, values_only=True)])
        self.assertContains(self.client.get(reverse('operations'), {'status': 'deleted'}), 'Исправлено')
        self.client.post(reverse('record_restore', args=['operations', item.pk]), {'revision': item.revision})
        item.refresh_from_db(); self.assertTrue(item.active)
        next_batch, _ = stage_import(legacy_bytes(income=900), 'next.xlsx')
        with self.assertRaises(ValidationError): commit_import(next_batch.pk)
        commit_import(next_batch.pk, replace_local_changes=True)
        response = self.client.post(reverse('record_restore', args=['operations', item.pk]), {'revision': item.revision}, follow=True)
        self.assertContains(response, 'архивному снимку')
        item.refresh_from_db(); self.assertFalse(item.active)

    def test_imported_delivery_edit_recalculates_export(self):
        batch, _ = stage_import(legacy_bytes(), 'first.xlsx'); commit_import(batch.pk)
        item = Delivery.objects.get()
        values = {'date': str(item.date), 'direction': 'out', 'partner': item.partner, 'vehicle': item.vehicle,
                  'gross': '100', 'tare': '20', 'discount': '1', 'price': '3500', 'notes': 'Коррекция', 'revision': 0}
        self.assertEqual(self.client.post(reverse('record_edit', args=['deliveries', item.pk]), values).status_code, 302)
        item.refresh_from_db()
        self.assertEqual(item.clean_weight, Decimal('79.200'))
        self.assertEqual(item.amount, Decimal('277200'))
        self.assertEqual(item.batch_id, batch.pk)
        book = load_workbook(BytesIO(self.client.get(reverse('export_download', args=['deliveries'])).content))
        self.assertEqual(book.active['M2'].value, 277200)

    def test_uzbek_pages_and_debt_validation(self):
        payment, _ = self.pay()
        self.client.cookies['django_language'] = 'uz'
        for name, args in [('partners', []), ('partner_detail', [self.partner.pk]), ('debt_detail', [self.debt.pk]),
                           ('debt_edit', [self.debt.pk]), ('debt_new', []), ('payment_new', [self.debt.pk]), ('payment_cancel', [payment.pk])]:
            response = self.client.get(reverse(name, args=args))
            self.assertContains(response, '<html lang="uz">')
            self.assertNotContains(response, 'Осталось погасить')
        self.assertContains(self.client.get(reverse('partners')), 'Qarzlar va hisob-kitoblar')
        self.assertContains(self.client.get(reverse('debt_detail', args=[self.debt.pk])), 'To‘lanishi qolgan')

    def test_auth_and_csrf_protect_all_debt_mutations(self):
        self.client.logout()
        self.assertEqual(self.client.post(reverse('payment_new', args=[self.debt.pk]), {}).status_code, 302)
        user = get_user_model().objects.create_user('debt-tester')
        secure = Client(enforce_csrf_checks=True); secure.force_login(user)
        self.assertEqual(secure.post(reverse('payment_new', args=[self.debt.pk]), {}).status_code, 403)
        self.assertEqual(secure.post(reverse('debt_delete', args=[self.debt.pk]), {}).status_code, 403)
        self.assertEqual(DebtPayment.objects.count(), 0)
