import hashlib
import os
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import Operation, Delivery, PartnerBalance, ImportBatch, SourceSheet, Workspace, WorkspaceMembership
from .forms import DeliveryForm
from .services.importing import stage_import, commit_import, inspect_workbook, OP_HEADERS, DEL_HEADERS
from .services.exporting import make_export


def book_bytes(book):
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def legacy_bytes(income=500):
    book = Workbook()
    cash = book.active
    cash.title = 'Приход-Расход'
    cash.append(['Дата', 'Расход сумма', 'Расход Наименование', 'Дата', 'Приход сумма', 'Приход Наименование'])
    cash['B3'] = '=SUM(B4:B5)'
    cash['E3'] = '=SUM(E4:E5)'
    cash.append([date(2026,9,1), 100, 'Металл', date(2026,9,1), income, 'Продажа'])
    cash.append([None, '=50*2', 'Транспорт', None, -25, 'Сторно'])
    main = book.create_sheet('тура ака')
    main.append(['Дата', 'Ойбек Пул берди', 'Тура ака Пул берди', 'Нима учун берди', 'Номер машини', 'Брутто', 'Тара', 'Нетто', '% скидка', 'Чистый', 'Цена', 'сумма Металл'])
    main['A4'] = date(2026,9,1)
    for cell, value in {'E4':'919','F4':26610,'G4':13890,'I4':1,'K4':3700,'H4':'=F4-G4','J4':'=H4-H4*I4/100','L4':'=J4*K4','L3':'=SUM(L4:L5)+1000'}.items():
        main[cell]=value
    summary = book.create_sheet('Кунлик')
    summary.append(['№','Партнёр','$','А ман карз','Ман Карз'])
    summary['B5']='Партнёр А';summary['D5']=1000;summary['E5']=250
    summary['B6']='Партнёр Б';summary['D6']=-500
    return book_bytes(book)


class WorkflowTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.settings_override = override_settings(MEDIA_ROOT=self.temp.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.addCleanup(self.temp.cleanup)
        self.user = get_user_model().objects.create_user('tester', password='strong-test-password')
        WorkspaceMembership.objects.create(user=self.user, workspace=Workspace.objects.order_by('pk').first())
        self.client.force_login(self.user)

    def test_legacy_reconciliation_and_date_inheritance(self):
        batch, _ = stage_import(legacy_bytes(), '01.09.2026.xlsx')
        self.assertEqual(batch.summary['income'], '475.00')
        self.assertEqual(batch.summary['expense'], '200.00')
        commit_import(batch.pk)
        self.assertEqual(Operation.objects.count(), 4)
        self.assertEqual(Operation.objects.filter(date=date(2026,9,1)).count(),4)
        delivery = Delivery.objects.get()
        self.assertEqual(delivery.clean_weight, Decimal('12592.800'))
        self.assertEqual(delivery.amount,Decimal('46593360.00'))
        self.assertEqual(PartnerBalance.objects.get(name='Партнёр Б').balance,-500)
        self.assertEqual(PartnerBalance.objects.get(name='Партнёр А').balance,750)

    def test_repeat_file_and_confirm_are_idempotent(self):
        data = legacy_bytes()
        batch, created = stage_import(data,'01.09.2026.xlsx')
        self.assertTrue(created)
        commit_import(batch.pk)
        duplicate, created = stage_import(data,'renamed.xlsx')
        self.assertFalse(created)
        self.assertEqual(duplicate.pk,batch.pk)
        with self.assertRaises(ValidationError):
            commit_import(batch.pk)
        self.assertEqual(Operation.objects.count(),4)

    def test_new_snapshot_archives_previous_but_preserves_manual(self):
        first,_=stage_import(legacy_bytes(),'01.09.2026.xlsx');commit_import(first.pk)
        manual=Operation.objects.create(date=date(2026,9,1),kind='income',amount=50,description='Вручную')
        second,_=stage_import(legacy_bytes(income=700),'02.09.2026.xlsx');commit_import(second.pk)
        first.refresh_from_db();manual.refresh_from_db()
        self.assertEqual(first.status,'superseded')
        self.assertFalse(Operation.objects.filter(batch=first,active=True).exists())
        self.assertEqual(Operation.objects.filter(active=True).count(),5)
        self.assertTrue(manual.active)
        self.assertTrue(Path(first.file.path).exists())

    def test_export_reimport_preserves_ids_and_signed_money(self):
        original=Operation.objects.create(date=date(2026,9,1),kind='income',amount=-25,description='=HYPERLINK("bad")')
        data=make_export('operations',Operation.objects.all())
        workbook=load_workbook(BytesIO(data),data_only=False)
        self.assertEqual(workbook.active['E2'].data_type,'s')
        self.assertEqual(workbook.active['D2'].value,-25)
        batch,_=stage_import(data,'operations.xlsx');commit_import(batch.pk)
        self.assertEqual(Operation.objects.count(),1)
        self.assertTrue(any('ID уже существует' in i['message'] for i in batch.issues))

    def test_invalid_template_row_blocks_entire_import(self):
        book=Workbook();book.active.title='Операции';book.active.append(OP_HEADERS)
        book.active.append([None,date(2026,9,1),'Приход',100,'Valid',None,None])
        book.active.append([None,date(2026,9,1),'Приход','oops','Invalid',None,None])
        batch,_=stage_import(book_bytes(book),'operations.xlsx')
        self.assertEqual(batch.summary['errors'],1)
        with self.assertRaises(ValidationError):commit_import(batch.pk)
        self.assertEqual(Operation.objects.count(),0)

    def test_bad_format_and_invalid_dates(self):
        with self.assertRaises(ValidationError):inspect_workbook(b'not-excel','test.xlsx')
        with self.assertRaises(ValidationError):inspect_workbook(legacy_bytes(),'test.xls')
        book=Workbook();book.active.title='Операции';book.active.append(OP_HEADERS)
        book.active.append([None,'31.02.2026','Приход',100,'Invalid date',None,None])
        batch,_=stage_import(book_bytes(book),'operations.xlsx')
        self.assertEqual(batch.summary['errors'],1)

    def test_broken_money_formula_blocks_snapshot(self):
        book=load_workbook(BytesIO(legacy_bytes()))
        book['Приход-Расход']['E4']='=#REF!'
        result=inspect_workbook(book_bytes(book),'01.09.2026.xlsx')
        self.assertGreater(result['summary']['errors'],0)

    def test_reordered_cash_columns_are_rejected(self):
        book=load_workbook(BytesIO(legacy_bytes()))
        book['Приход-Расход']['E1']='Расход сумма'
        with self.assertRaises(ValidationError):inspect_workbook(book_bytes(book),'01.09.2026.xlsx')

    def test_missing_dates_not_fabricated_from_other_cash_column(self):
        book=load_workbook(BytesIO(legacy_bytes()))
        book['Приход-Расход']['D4']=None
        batch,_=stage_import(book_bytes(book),'01.09.2026.xlsx')
        self.assertEqual(batch.summary['undated'],2)

    def test_delivery_validation_and_live_totals(self):
        values={'date':'2026-09-01','direction':'out','partner':'Test','vehicle':'919','gross':'100','tare':'101','discount':'2','price':'3500'}
        form=DeliveryForm(values)
        self.assertFalse(form.is_valid());self.assertIn('tare',form.errors)
        values['tare']='20'
        form=DeliveryForm(values);self.assertTrue(form.is_valid(),form.errors)
        item=form.save();self.assertEqual(item.amount,Decimal('274400.00'))
        item.price=Decimal('4000');item.save();self.assertEqual(item.amount,Decimal('313600.00'))

    def test_manual_operation_http_lifecycle(self):
        values={'date':'2026-09-01','kind':'income','amount':'100','description':'Ручная операция','category':'Прочее','partner':''}
        response=self.client.post(reverse('operation_new'),values)
        self.assertRedirects(response,reverse('record_detail',args=['operations',Operation.objects.get(description='Ручная операция').pk]))
        item=Operation.objects.get(description='Ручная операция')
        values['amount']='125.50'
        values['revision']=item.revision
        self.assertEqual(self.client.post(reverse('record_edit',args=['operations',item.pk]),values).status_code,302)
        item.refresh_from_db();self.assertEqual(item.amount,Decimal('125.50'))
        values['amount']='0'
        response=self.client.post(reverse('operation_new'),values)
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'Сумма операции не может быть нулевой')
        self.assertEqual(self.client.post(reverse('record_delete',args=['operations',item.pk]), {'revision': item.revision}).status_code,302)
        item.refresh_from_db();self.assertFalse(item.active)

    def test_pages_filters_and_downloads(self):
        batch,_=stage_import(legacy_bytes(),'01.09.2026.xlsx');commit_import(batch.pk)
        for route in ['dashboard','operations','deliveries','partners','imports','exports','sheets','help']:
            response=self.client.get(reverse(route),follow=True);self.assertEqual(response.status_code,200,(route,response.content[:500]))
        for sheet in SourceSheet.objects.all():
            self.assertEqual(self.client.get(reverse('sheet_detail',args=[sheet.pk])).status_code,200)
        self.assertEqual(self.client.get(reverse('import_detail',args=[batch.pk])).status_code,200)
        for kind in ['operations','deliveries','partners','report']:
            response=self.client.get(reverse('export_download',args=[kind]))
            self.assertEqual(response.status_code,200)
            book=load_workbook(BytesIO(response.content));self.assertGreaterEqual(len(book.sheetnames),1)
        response=self.client.get(reverse('export_download',args=['operations']),{'start':'2027-01-01'})
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row,1)
        self.assertEqual(self.client.get(reverse('export_download',args=['operations']),{'start':'invalid'}).status_code,400)
        response=self.client.get(reverse('operations'),{'q':'Сторно'})
        self.assertEqual(response.context['page'].paginator.count,1)

    def test_auth_csrf_and_imported_record_access(self):
        anonymous=Client()
        self.assertEqual(anonymous.get('/').status_code,302)
        batch,_=stage_import(legacy_bytes(),'01.09.2026.xlsx');commit_import(batch.pk)
        source=anonymous.get(reverse('source_download',args=[batch.pk]));self.assertEqual(source.status_code,302)
        secure=Client(enforce_csrf_checks=True);secure.force_login(self.user)
        self.assertEqual(secure.post(reverse('import_confirm',args=[batch.pk])).status_code,403)
        item=Operation.objects.first()
        self.assertEqual(self.client.get(reverse('record_edit',args=['operations',item.pk])).status_code,200)
        self.assertEqual(secure.post(reverse('record_delete',args=['operations',item.pk]), {'revision': 0}).status_code,403)

    def test_templates_and_delivery_round_trip(self):
        for kind,title,headers in [('operations','Операции',OP_HEADERS),('deliveries','Поставки',DEL_HEADERS)]:
            data=make_export(kind,template=True)
            book=load_workbook(BytesIO(data));self.assertEqual(book.active.title,title)
            self.assertEqual([c.value for c in book.active[1]],headers)
        item=Delivery.objects.create(date=date(2026,9,1),direction='in',partner='Феруз',vehicle='919',gross=18800,tare=13890,discount=1,price=3500)
        data=make_export('deliveries',deliveries=Delivery.objects.all())
        result=inspect_workbook(data,'deliveries.xlsx')
        self.assertEqual(result['summary']['errors'],0)
        self.assertEqual(result['summary']['deliveries'],0)


class SuppliedWorkbookTests(TestCase):
    def test_real_source_totals(self):
        workbook_path = os.getenv('WORKBOOK_PATH')
        if not workbook_path:
            self.skipTest('Set WORKBOOK_PATH for source workbook reconciliation')
        path = Path(workbook_path)
        if not path.is_file():
            self.skipTest('WORKBOOK_PATH must point to an existing workbook')
        data=path.read_bytes()
        before=hashlib.sha256(data).hexdigest()
        result=inspect_workbook(data,path.name)
        book=load_workbook(BytesIO(data),data_only=True)
        for key,cell in [('income','E3'),('expense','B3')]:
            self.assertAlmostEqual(Decimal(result['summary'][key]),Decimal(str(book['Приход-Расход'][cell].value)),delta=Decimal('.05'))
        out_weight=sum(Decimal(d['clean_weight']) for d in result['deliveries'] if d['direction']=='out')
        self.assertAlmostEqual(out_weight,Decimal(str(book['тура ака']['J3'].value)),delta=Decimal('.001'))
        net=sum(Decimal(p['balance']) for p in result['balances'])
        self.assertAlmostEqual(net,Decimal(str(book['Кунлик']['G2'].value)),delta=Decimal('.05'))
        self.assertTrue(any(i['cell']=='H3' and '#REF!' in i['message'] for i in result['issues']))
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
